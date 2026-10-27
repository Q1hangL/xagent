# The repository's Ruff and isort configurations classify the MCP package
# differently; keep this module's import grouping stable for both hooks.
# isort: skip_file

import errno
import io
import ipaddress
import json
import logging
import os
import re
import socket
import threading
import time
import uuid
from collections import Counter
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import urljoin, urlsplit, urlunsplit

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build  # type: ignore[import-not-found]
from googleapiclient.http import MediaIoBaseUpload  # type: ignore[import-not-found]
from pydantic import BeforeValidator

from mcp.server.fastmcp import FastMCP

from ....core.utils import public_requests as requests
from .utils import (
    GoogleFileKind,
    allowed_dirs_from_env,
    google_file_error_message,
    resolve_google_file_id,
    setup_proxy_env,
)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("google-slides-mcp")

# Ensure standard proxy environment variables are set to prevent hanging requests
setup_proxy_env()

mcp = FastMCP("google-slides-mcp")

# "(?!e/)" skips a published-to-the-web link (".../d/e/<publish-id>/pub"),
# whose "e" segment is not the presentation id.
_PRESENTATION_URL_ID_PATTERN = re.compile(
    r"/presentation/(?:u/\d+/)?d/(?!e/)([a-zA-Z0-9_-]+)"
)
_PRESENTATION_KIND = GoogleFileKind(
    product="Google Slides",
    noun="presentation",
    link_example="https://docs.google.com/presentation/d/...",
    create_tool="google_slides_create_presentation",
    office_file="a PowerPoint file (.pptx)",
    office_reader="read_pptx",
)
_PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
_GOOGLE_SLIDES_MIME_TYPE = "application/vnd.google-apps.presentation"
_PPTX_UPLOAD_ALLOWED_DIRS_ENV_VAR = "XAGENT_GOOGLE_DRIVE_FILE_ALLOWED_DIRS"
# Google Drive's PPTX-to-Slides conversion rejects files over 100 MB.  Check
# the local file before parsing it with python-pptx or starting an upload.
_MAX_PPTX_UPLOAD_BYTES = 100_000_000

# Keep track of default pages created by this MCP process.  A presentation can
# legitimately contain an intentional blank page, so add_slide must not infer
# that every sole empty page is Google's implementation detail.  The returned
# id can also be passed explicitly for calls that cross process boundaries.
_CREATED_DEFAULT_SLIDES: dict[str, str] = {}
_PRESERVED_BLANK_SLIDES: dict[str, set[str]] = {}
_MAX_TRACKED_DEFAULT_SLIDES = 256


def _normalize_layout(value: object) -> object:
    """Case/whitespace-normalize `layout` before FastMCP's Pydantic layer
    validates it against the Literal below. Without this, an MCP tool
    call (the only way a real agent ever invokes this tool — direct
    Python calls, used by this file's own tests, bypass FastMCP
    validation entirely) with e.g. "title_and_body" would be rejected by
    Pydantic's case-sensitive Literal check *before* the function body's
    own normalization ever ran, making that normalization dead code for
    every real caller."""
    return value.strip().upper() if isinstance(value, str) else value


_Layout = Annotated[
    Literal["TITLE", "TITLE_AND_BODY", "TITLE_ONLY", "SECTION_HEADER"],
    BeforeValidator(_normalize_layout),
]

# Predefined layouts we know how to fill in, mapped to the (title_placeholder,
# body_placeholder) types Slides creates for each one. `None` means that slot
# doesn't exist on the layout at all. This table only records that factual
# placeholder shape; `is_bulleted`/`body_required` policy is derived from it
# once in google_slides_add_slide (both happen to be True only for
# TITLE_AND_BODY today, i.e. exactly when body_placeholder == "BODY").
#
# Deliberately limited to the layouts whose placeholder composition is well
# established (matching Google's official Apps Script PredefinedLayout docs
# and this file's pre-existing TITLE_AND_BODY mapping). Google's docs
# explicitly warn a predefined layout's placeholder set "may have been
# changed" and don't enumerate it for every layout, so less-common ones
# (SECTION_TITLE_AND_DESCRIPTION, ONE_COLUMN_TEXT, MAIN_POINT, BIG_NUMBER)
# are intentionally left out rather than guessed at — use
# google_slides_batch_update for those until verified against a live API.
_LAYOUT_PLACEHOLDERS: dict[str, tuple[str | None, str | None]] = {
    "TITLE": ("CENTERED_TITLE", "SUBTITLE"),
    "TITLE_AND_BODY": ("TITLE", "BODY"),
    "TITLE_ONLY": ("TITLE", None),
    "SECTION_HEADER": ("TITLE", None),
}

# Leading "•"/"-"/"*" marker at the start of a line, with or without a
# following space (so a marker glued to a word, e.g. "•First", or with
# nothing after it at all, e.g. a bare "•" on its own line, is still
# recognized). Captures any leading Unicode whitespace so it can be
# converted to nesting-level tabs by _leading_whitespace_to_tabs below.
#
# "-" and "*" also have common non-bullet meanings, so the "glued, no
# space" case is deliberately narrower for them than for "•" (which has no
# other meaning in plain text):
#   - "-" is only glued-recognized when followed directly by a letter
#     ("-Nospace"); anything else — a digit, another "-", a currency
#     symbol, a decimal point — is left alone, so a negative number's
#     sign ("-5%", "-$5m", "-.5%") or an em-dash/en-dash convention
#     ("-- Author", "-– Author") is never silently eaten.
#   - "*" is only treated as a marker when followed by real whitespace
#     ("* item") or nothing; a bare glued "*" ("*emphasis* text") is left
#     alone so markdown-style emphasis isn't corrupted.
#
# The "glued to a letter" lookahead is Unicode-aware (Python's default
# \w), so e.g. "-中文条目" is recognized and stripped the same way
# "-Nospace" is — intentional, not an ASCII-only oversight.
#
# Only "•"/"-"/"*" are recognized. Other list conventions (numbered lists
# like "1.", or other dash/arrow glyphs like "–"/"—"/"→") are left as
# literal text and will visually double up with Slides' own bullet glyph
# once createParagraphBullets is applied.
_BULLET_PREFIX_PATTERN = re.compile(
    r"^(\s*)(?:"
    r"•(?:[ \t]+|(?=\S)|$)"
    r"|-(?:[ \t]+|(?=[^\W\d_])|$)"
    r"|\*(?:[ \t]+|$)"
    r")"
)


def _leading_whitespace_to_tabs(leading: str) -> str:
    """Convert captured leading whitespace into the literal tab
    characters Slides' createParagraphBullets actually uses to infer
    nesting level (plain spaces are otherwise inserted as literal,
    non-nesting text). Each literal tab in the input counts as one level;
    otherwise every 2 spaces — the typical hand-typed nested-bullet
    convention — counts as one level, with any non-empty indentation
    counting as at least one level. Mixing tabs and spaces in the same
    run is not combined: if any tab is present, only the tabs are
    counted and any spaces in that same leading run are ignored."""
    if not leading:
        return ""
    level = leading.count("\t") or max(1, len(leading) // 2)
    return "\t" * level


# Upper bound on how many leading markers _strip_line unwraps from one
# line. Real content never nests anywhere near this deep (the case this
# exists for, "• - nested", needs 2); the cap exists purely to bound the
# work: each iteration re-scans the line from its new start, so without a
# cap a degenerate line of N glued markers (e.g. "•" * 200_000, from a
# runaway/malformed LLM completion — body has no length limit) makes
# _strip_line O(n^2), measured to take over a second at N=200_000.
_MAX_MARKER_STRIPS_PER_LINE = 8


def _strip_bullet_prefixes(text: str) -> str:
    """Drop a leading "•"/"-"/"*" marker from each line — repeatedly, up
    to _MAX_MARKER_STRIPS_PER_LINE times, so a line with more than one
    leading marker (e.g. "• - nested") is fully unwrapped rather than
    leaving a residual marker as literal text — and convert any leading
    indentation into nesting-level tabs. See _BULLET_PREFIX_PATTERN for
    the narrower rules that keep this from corrupting non-bullet content.

    The public docstrings tell callers not to type literal bullet glyphs,
    but callers may do so anyway; once we ask Slides to render real
    bulleted paragraphs (see createParagraphBullets below) those glyphs
    would double up with Slides' own bullet, so strip them first.
    """

    def _strip_line(line: str) -> str:
        for _ in range(_MAX_MARKER_STRIPS_PER_LINE):
            match = _BULLET_PREFIX_PATTERN.match(line)
            if match is None:
                return line
            stripped = _leading_whitespace_to_tabs(match.group(1)) + line[match.end() :]
            if stripped == line:
                return line
            line = stripped
        return line

    return "\n".join(_strip_line(line) for line in text.split("\n"))


def _error(message: str) -> str:
    return json.dumps({"status": "error", "message": message}, ensure_ascii=False)


def _presentation_error(
    exc: Exception,
    presentation_id: str,
    *,
    editing: bool = False,
    opened: bool = False,
) -> str:
    """``presentation_id`` is the tool's own argument, before it was resolved.

    Pass ``opened=True`` once the tool has read the presentation: it is then
    a Google Slides presentation, whatever link named it, so a 400 for a
    later request is about that request and keeps Google's own error."""
    return _error(
        google_file_error_message(
            exc,
            _PRESENTATION_KIND,
            editing=editing,
            file_link_or_id=None if opened else presentation_id,
        )
    )


def _delete_imported_presentation(drive_service: Any, presentation_id: str) -> bool:
    """Delete a Drive presentation that failed import validation."""
    try:
        (
            drive_service.files()
            .delete(fileId=presentation_id, supportsAllDrives=True)
            .execute()
        )
        return True
    except Exception as exc:  # noqa: BLE001
        logger.warning(
            "Could not clean up failed imported presentation %s: %s",
            presentation_id,
            exc,
        )
        return False


def _normalize_newlines(text: str) -> str:
    """Normalize CRLF/lone-CR line endings to plain "\\n", so text pasted
    from a Windows editor doesn't leave stray "\\r" characters embedded in
    what gets inserted into Slides."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _normalize_title(title: str) -> str:
    """Normalize line endings, then collapse any embedded newline (and
    surrounding whitespace) into a single space — a title is expected to
    be a single line, so this avoids silently producing a multi-paragraph
    title placeholder."""
    return re.sub(r"\s*\n\s*", " ", _normalize_newlines(title))


def _prepare_body_text(body: str, is_bulleted: bool) -> str:
    """Normalize body text before it's inserted: strip literal bullet
    markers (see _strip_bullet_prefixes) when the target placeholder
    bullets its content, then drop any blank lines (from a marker that
    stripped down to nothing, blank lines already in the input, or a
    trailing newline) so no empty paragraph survives — for bulleted
    content this avoids createParagraphBullets putting a bullet glyph on
    an empty paragraph (a visibly floating bullet point); for
    non-bulleted content (e.g. a TITLE layout's SUBTITLE) it avoids stray
    blank lines in the placeholder.

    Must run before any "is body non-empty" check, not just before
    insertion — a marker-only or blank-only body (e.g. "•   ") would
    otherwise pass a raw truthiness/strip check and only turn out empty
    right before the API call, silently producing a content-less slide.
    """
    if is_bulleted:
        body = _strip_bullet_prefixes(body)
    return "\n".join(line for line in body.split("\n") if line.strip())


def _body_insert_requests(
    body_id: str, body: str, is_bulleted: bool
) -> list[dict[str, Any]]:
    """insertText for `body_id` — `body` must already be prepared via
    _prepare_body_text — plus createParagraphBullets when the target
    placeholder is a real bulleted-list BODY (not e.g. a plain SUBTITLE)."""
    requests: list[dict[str, Any]] = [
        {"insertText": {"objectId": body_id, "text": body}}
    ]
    if is_bulleted:
        requests.append(
            {
                "createParagraphBullets": {
                    "objectId": body_id,
                    "textRange": {"type": "ALL"},
                    "bulletPreset": "BULLET_DISC_CIRCLE_SQUARE",
                }
            }
        )
    return requests


# Placeholder types that fill the "title" vs "body" role of a slide, used
# to locate an existing slide's placeholders by type when editing it (as
# opposed to _LAYOUT_PLACEHOLDERS above, which picks placeholder types
# when *creating* a slide from a predefined layout). BODY covers this
# file's own generated slides; OBJECT is included too since it's the
# common "body-like" placeholder type on slides from an imported/
# non-standard-theme presentation (e.g. one converted from PowerPoint),
# which google_slides_update_slide may be asked to edit even though
# google_slides_add_slide never creates that type itself.
_TITLE_PLACEHOLDER_TYPES = {"TITLE", "CENTERED_TITLE"}
_BODY_PLACEHOLDER_TYPES = {"BODY", "SUBTITLE", "OBJECT"}


def _find_slide(presentation: dict[str, Any], slide_id: str) -> dict[str, Any] | None:
    return next(
        (
            slide
            for slide in presentation.get("slides", [])
            if slide.get("objectId") == slide_id
        ),
        None,
    )


def _find_placeholders(
    slide: dict[str, Any],
) -> dict[str, tuple[dict[str, Any], str]]:
    """Map "title"/"body" to (element, placeholder_type) for a slide's
    shapes, so callers can target the right shape without knowing the
    layout-specific ids assigned when the slide was created.

    If a slide has more than one placeholder of the same role (not
    reachable via this file's own google_slides_add_slide, which only
    ever creates one of each, but possible for a slide created some other
    way), only the first one encountered is kept — there's no way to
    disambiguate further from the role alone.
    """
    found: dict[str, tuple[dict[str, Any], str]] = {}
    for element in slide.get("pageElements", []):
        placeholder = element.get("shape", {}).get("placeholder")
        if not placeholder:
            continue
        placeholder_type = placeholder.get("type", "")
        if placeholder_type in _TITLE_PLACEHOLDER_TYPES:
            found.setdefault("title", (element, placeholder_type))
        elif placeholder_type in _BODY_PLACEHOLDER_TYPES:
            found.setdefault("body", (element, placeholder_type))
    return found


def _get_google_credentials() -> Credentials:
    token = os.environ.get("GOOGLE_ACCESS_TOKEN")
    refresh_token = os.environ.get("GOOGLE_REFRESH_TOKEN")
    client_id = os.environ.get("GOOGLE_CLIENT_ID")
    client_secret = os.environ.get("GOOGLE_CLIENT_SECRET")

    if not token:
        raise ValueError("GOOGLE_ACCESS_TOKEN environment variable is missing")

    creds_kwargs = {"token": token}
    if refresh_token and client_id and client_secret:
        creds_kwargs.update(
            {
                "refresh_token": refresh_token,
                "client_id": client_id,
                "client_secret": client_secret,
                "token_uri": "https://oauth2.googleapis.com/token",
            }
        )

    credentials = Credentials(**creds_kwargs)
    return credentials


def get_slides_service() -> Any:
    credentials = _get_google_credentials()
    return build("slides", "v1", credentials=credentials)


def get_drive_service() -> Any:
    credentials = _get_google_credentials()
    return build("drive", "v3", credentials=credentials)


def _resolve_pptx_upload_path(file_path: str) -> Path:
    """Resolve a generated PPTX and keep uploads inside approved roots."""
    try:
        allowed_dirs = allowed_dirs_from_env(_PPTX_UPLOAD_ALLOWED_DIRS_ENV_VAR)
    except ValueError:
        logger.warning("Invalid PPTX upload directory configuration")
        raise ValueError("Upload directory configuration is invalid") from None

    requested_path = Path(file_path).expanduser()
    candidates = (
        [requested_path]
        if requested_path.is_absolute()
        else [Path.cwd() / requested_path]
        + [directory / requested_path for directory in allowed_dirs]
    )
    local_path: Path | None = None
    authorized_candidate: Path | None = None
    path_error: OSError | None = None
    for candidate in candidates:
        try:
            resolved_candidate = candidate.resolve()
        except RuntimeError:
            # Python 3.11/3.12 raises RuntimeError for symlink loops. Keep the
            # original path so stat() can surface the expected ELOOP error.
            resolved_candidate = candidate
        except OSError as exc:
            logger.warning("Could not resolve PPTX path %s: %s", candidate, exc)
            raise ValueError("file path could not be resolved safely") from None
        if not any(
            resolved_candidate.is_relative_to(directory) for directory in allowed_dirs
        ):
            continue
        authorized_candidate = authorized_candidate or resolved_candidate
        try:
            resolved_candidate.stat()
        except OSError as exc:
            if exc.errno != errno.ENOENT:
                logger.warning(
                    "Could not inspect PPTX path %s: %s", resolved_candidate, exc
                )
                path_error = path_error or exc
            continue
        if resolved_candidate.is_file():
            local_path = resolved_candidate
            break

    if local_path is None:
        if authorized_candidate is None:
            try:
                rejected_path = candidates[0].resolve()
            except (OSError, RuntimeError):
                rejected_path = candidates[0]
            logger.warning(
                "Rejected PPTX path %s outside allowed directories: %s",
                rejected_path,
                ", ".join(str(directory) for directory in allowed_dirs),
            )
            raise PermissionError(
                "file path is outside the allowed directories; provide a PPTX "
                "inside the task workspace"
            )
        if path_error is not None:
            raise path_error
        raise FileNotFoundError(f"File not found: {file_path}")

    if local_path.suffix.lower() != ".pptx":
        raise ValueError("file_path must point to a .pptx file")
    file_size = local_path.stat().st_size
    if file_size == 0:
        raise ValueError("PPTX file is empty")
    if file_size > _MAX_PPTX_UPLOAD_BYTES:
        raise ValueError("PPTX file exceeds the 100 MB Google Slides conversion limit")
    return local_path


def _normalize_comparison_text(value: str) -> str:
    """Normalize whitespace for PPTX/Slides text-preservation checks."""
    return re.sub(r"\s+", " ", value).strip()


def _shape_texts(shape: Any) -> list[str]:
    """Extract text from a python-pptx shape, including grouped shapes."""
    texts: list[str] = []
    text = getattr(shape, "text", "")
    if text and str(text).strip():
        texts.append(str(text))
    grouped_shapes = getattr(shape, "shapes", None)
    if grouped_shapes is not None:
        for child in grouped_shapes:
            texts.extend(_shape_texts(child))
    return texts


def _shape_has_meaningful_content(shape: Any) -> bool:
    """Return whether a PPTX shape represents content beyond an empty placeholder."""
    text = getattr(shape, "text", "")
    if text and str(text).strip():
        return True

    grouped_shapes = getattr(shape, "shapes", None)
    if grouped_shapes is not None:
        return any(_shape_has_meaningful_content(child) for child in grouped_shapes)

    return not bool(getattr(shape, "is_placeholder", False))


def _extract_pptx_slide_content(file_path: Path) -> list[tuple[str, bool]]:
    """Return normalized text and meaningful-content flags for each PPTX slide."""
    from pptx import Presentation

    presentation = Presentation(str(file_path))
    slide_content: list[tuple[str, bool]] = []
    for slide in presentation.slides:
        texts = [text for shape in slide.shapes for text in _shape_texts(shape)]
        slide_content.append(
            (
                _normalize_comparison_text(" ".join(texts)),
                any(_shape_has_meaningful_content(shape) for shape in slide.shapes),
            )
        )
    return slide_content


def _extract_pptx_slide_text(file_path: Path) -> list[str]:
    """Return normalized visible text for each local PPTX slide."""
    return [text for text, _has_content in _extract_pptx_slide_content(file_path)]


def _record_created_default_slide(presentation_id: str, slide_id: str) -> None:
    if len(_CREATED_DEFAULT_SLIDES) >= _MAX_TRACKED_DEFAULT_SLIDES:
        _CREATED_DEFAULT_SLIDES.pop(next(iter(_CREATED_DEFAULT_SLIDES)))
    _CREATED_DEFAULT_SLIDES[presentation_id] = slide_id


def _resolve_presentation_id(presentation_id: str, *, editing: bool = False) -> str:
    """Accept either a bare presentation id or a full Google Slides URL.

    Pass ``editing=True`` from a tool that changes the presentation."""
    return resolve_google_file_id(
        presentation_id,
        _PRESENTATION_URL_ID_PATTERN,
        "presentation_id",
        _PRESENTATION_KIND,
        editing=editing,
    )


def _element_text(element: dict[str, Any]) -> str:
    text_elements = element.get("shape", {}).get("text", {}).get("textElements", [])
    own_text = "".join(
        text_element.get("textRun", {}).get("content", "")
        or text_element.get("autoText", {}).get("content", "")
        for text_element in text_elements
    )
    grouped_text = "".join(
        _element_text(child)
        for child in element.get("elementGroup", {}).get("children", [])
    )
    return own_text + grouped_text


def _text_tokens(value: str) -> Counter[str]:
    """Return order-independent visible-text tokens for import validation."""
    return Counter(re.findall(r"\w+", value.casefold(), re.UNICODE))


def _contains_text_tokens(expected: str, actual: str) -> bool:
    """Return whether ``actual`` preserves all tokens from ``expected``.

    Drive may reorder page elements while converting a PPTX.  Comparing the
    token multisets keeps the validation useful for detecting dropped detail
    text without rejecting an otherwise correct conversion solely because a
    title and body shape were returned in a different order.
    """
    expected_tokens = _text_tokens(expected)
    if not expected_tokens:
        return expected.casefold() in actual.casefold()
    return not (expected_tokens - _text_tokens(actual))


def _slide_summary(slide: dict[str, Any], index: int) -> dict[str, Any]:
    texts = [
        text
        for element in slide.get("pageElements", [])
        if (text := _element_text(element).strip())
    ]
    return {
        "slide_number": index + 1,
        "object_id": slide.get("objectId"),
        "text": texts,
    }


def _slide_is_empty(slide: dict[str, Any]) -> bool:
    """Return whether a slide has no meaningful page content.

    A visual slide may intentionally contain only an image, chart, line, or
    other non-text element. Treating "no text" as empty would reject those
    legitimate slides after PPTX import. Empty placeholder shapes do not count
    as content, but any non-placeholder shape or non-shape page element does.
    """
    page_elements = slide.get("pageElements", [])
    if not page_elements:
        return True
    for element in page_elements:
        shape = element.get("shape")
        if shape is None:
            return False
        if _element_text(element).strip():
            return False
        if "placeholder" not in shape:
            return False
    return True


def _slide_is_untouched_default(slide: dict[str, Any]) -> bool:
    """Return whether a slide is safe to infer as Google's initial page.

    The stateless fallback has no object id to prove provenance.  Restrict it
    to a truly element-free slide with no custom background; empty title
    placeholders and background-only covers are user content and must remain.
    """
    if slide.get("pageElements"):
        return False
    page_properties = slide.get("pageProperties") or {}
    background = page_properties.get("pageBackgroundFill") or {}
    return background.get("propertyState", "INHERIT") == "INHERIT" and not any(
        key in background for key in ("solidFill", "stretchedPictureFill")
    )


def _title_placeholder_id(slide: dict[str, Any]) -> str | None:
    """Return the object id of the slide's first title placeholder."""
    for element in slide.get("pageElements", []):
        placeholder = (element.get("shape") or {}).get("placeholder") or {}
        if placeholder.get("type") in _TITLE_PLACEHOLDER_TYPES:
            object_id = element.get("objectId")
            if isinstance(object_id, str) and object_id:
                return object_id
    return None


def _slide_is_titled_default(slide: dict[str, Any], deck_title: object) -> bool:
    """Return whether a slide is the default page as
    google_slides_create_presentation leaves it: placeholders only, with the
    deck title in its title placeholder and every other placeholder empty.

    Only used for a page already known to be the default (by its id), so
    that writing the deck title into it at create time does not make it
    look like user content to google_slides_add_slide. The comparison uses
    the deck's current title: if the deck was renamed after create, the page
    shows the old title and is kept like a page the user edited.
    """
    if not isinstance(deck_title, str):
        return False
    expected = _normalize_title(deck_title).strip()
    if not expected:
        return False
    titled = False
    for element in slide.get("pageElements", []):
        shape = element.get("shape")
        if shape is None or "placeholder" not in shape:
            return False
        text = _element_text(element).strip()
        if not text:
            continue
        if (
            not titled
            and shape["placeholder"].get("type") in _TITLE_PLACEHOLDER_TYPES
            and text == expected
        ):
            titled = True
            continue
        return False
    return titled


@mcp.tool()
def google_slides_get_presentation(presentation_id: str) -> str:
    """
    Read a Google Slides presentation by id or full URL.
    Returns the title and the text content of each slide.

    Presentations are opened only by link or id; this connector cannot
    search for or list presentations by name. When the user names a presentation
    without giving its link, use google_drive_search to find its id if
    that tool is available (with per-file Drive access it only finds
    files created through this app or granted to it); otherwise, or if
    it finds nothing, ask the user to paste the link
    (https://docs.google.com/presentation/d/...). Connecting Google Drive is
    not needed to open a presentation by its link.
    """
    try:
        pres_id = _resolve_presentation_id(presentation_id)
        service = get_slides_service()
        presentation = service.presentations().get(presentationId=pres_id).execute()

        slides = [
            _slide_summary(slide, index)
            for index, slide in enumerate(presentation.get("slides", []))
        ]
        return json.dumps(
            {
                "status": "success",
                "presentation_id": presentation.get("presentationId"),
                "title": presentation.get("title"),
                "slide_count": len(slides),
                "slides": slides,
            },
            ensure_ascii=False,
        )
    except Exception as e:
        logger.error(f"Error getting presentation: {e}")
        return _presentation_error(e, presentation_id)


@mcp.tool()
def google_slides_create_presentation(title: str) -> str:
    """
    Create a new Google Slides presentation with the given title.

    This is a plain text/layout primitive. For a new deck that needs visual
    design, create a styled PPTX first and call google_slides_import_pptx;
    do not build the deck by repeatedly calling google_slides_add_slide.

    Google creates a default, empty slide as part of this operation. When
    that page has a title placeholder, this call writes the deck title into
    it, so the deck opens on a title cover even if the page is never
    replaced. Keep the page until the first content slide is created:
    google_slides_add_slide removes it in the same batchUpdate after creating
    the real slide, so the API never has to delete the only page from the
    presentation. The response reports that page as default_slide_id (null
    when there is no such page); pass it to the first google_slides_add_slide
    call, because each call may run in a separate process that cannot
    recognize the default page on its own. Unless the user asked for a
    different first slide, make that first slide the deck's cover:
    layout="TITLE" with the deck title (and an optional subtitle in body).
    Without the id the default page stays first, so do not add a second
    cover then. The response's next_step repeats this; it does not ask for
    slides the user did not request.
    """
    try:
        service = get_slides_service()
        presentation = service.presentations().create(body={"title": title}).execute()
        pres_id = presentation.get("presentationId")
        if not pres_id:
            raise ValueError("Google Slides did not return a presentation_id")

        # The create response normally includes the initial page, but Google
        # does not document that shape as a stable part of the create
        # response. Fall back to a read when it is omitted so the cleanup is
        # safe across client/API versions.
        slides = presentation.get("slides")
        if slides is None:
            created = service.presentations().get(presentationId=pres_id).execute()
            slides = created.get("slides", [])

        default_slide_id = None
        default_page_titled = False
        if len(slides) == 1 and _slide_is_empty(slides[0]):
            default_slide_id = slides[0].get("objectId")
            if not default_slide_id:
                raise ValueError(
                    "Google Slides returned an empty initial slide without an object_id"
                )
            _record_created_default_slide(pres_id, default_slide_id)
            # Without default_slide_id, google_slides_add_slide keeps a
            # default page that has placeholders, so a caller that does not
            # pass the id would otherwise leave a blank first slide. With
            # the deck title in it, that page is a title cover instead. A
            # failed write leaves the page blank, as before.
            title_placeholder_id = _title_placeholder_id(slides[0])
            cover_title = _normalize_title(presentation.get("title") or title).strip()
            if title_placeholder_id and cover_title:
                try:
                    service.presentations().batchUpdate(
                        presentationId=pres_id,
                        body={
                            "requests": [
                                {
                                    "insertText": {
                                        "objectId": title_placeholder_id,
                                        "text": cover_title,
                                    }
                                }
                            ]
                        },
                    ).execute()
                    default_page_titled = True
                except Exception as exc:  # noqa: BLE001
                    logger.warning(
                        "Could not write the deck title into the default page "
                        "of %s: %s",
                        pres_id,
                        exc,
                    )

        response: dict[str, Any] = {
            "status": "success",
            "presentation_id": pres_id,
            "title": presentation.get("title"),
            "link": f"https://docs.google.com/presentation/d/{pres_id}/edit",
            "slide_count": len(slides),
            "default_slide_id": default_slide_id,
        }
        if default_slide_id:
            first_call = (
                "When you add slides to this presentation, pass "
                f"default_slide_id={default_slide_id!r} on the first "
                f"google_slides_add_slide call for presentation_id={pres_id!r}. "
            )
            cover = (
                "the cover: layout='TITLE' with the deck title as title and an "
                "optional subtitle as body."
            )
            if default_page_titled:
                response["next_step"] = (
                    f"{first_call}The id lets that call replace the default "
                    "page, which already shows the deck title. Unless the user "
                    f"asked for a different first slide, make that call {cover} "
                    "Without the id the default page stays first as a "
                    "title-only cover, so do not add a second cover; a "
                    "different first slide the user asked for would then come "
                    "second."
                )
            else:
                response["next_step"] = (
                    f"{first_call}The id lets that call replace Google's blank "
                    "default page; without it the page can remain as an empty "
                    "first slide. Unless the user asked for a different first "
                    f"slide, make it {cover}"
                )
        return json.dumps(response, ensure_ascii=False)
    except Exception as e:
        logger.error(f"Error creating presentation: {e}")
        return _error(str(e))


@mcp.tool()
def google_slides_import_pptx(file_path: str, title: str = "") -> str:
    """Import a designed local PPTX as a native Google Slides presentation.

    Use this for requests that ask for a beautiful, polished, or visually
    designed deck. The existing PPTX generator owns layout and styling; Drive
    converts its bytes to a native Google Slides file, preserving editable
    slides instead of falling back to raw title/body placeholders.

    ``file_path`` must point to a non-empty ``.pptx`` inside the current task
    workspace (or another configured Google Drive upload directory). The
    conversion is read back through the Slides API before this tool reports
    success, and an unexpected blank slide or dropped visual content is
    returned as ``validation_failed`` rather than being presented as a
    finished deck.
    """
    drive_service: Any | None = None
    presentation_id: str | None = None
    try:
        local_path = _resolve_pptx_upload_path(file_path)
        resolved_title = title.strip() or local_path.stem
        if not resolved_title:
            return _error("title cannot be empty")
        try:
            expected_slide_content = _extract_pptx_slide_content(local_path)
            expected_slide_text = [
                text for text, _has_content in expected_slide_content
            ]
            expected_slide_has_content = [
                has_content for _text, has_content in expected_slide_content
            ]
        except Exception as exc:  # noqa: BLE001
            logger.error("Could not inspect PPTX text before import: %s", exc)
            return _error("Could not inspect PPTX text before import")

        drive_service = get_drive_service()
        with local_path.open("rb") as file_handle:
            media = MediaIoBaseUpload(
                file_handle, mimetype=_PPTX_MIME_TYPE, resumable=True
            )
            imported_file = (
                drive_service.files()
                .create(
                    body={
                        "name": resolved_title,
                        "mimeType": _GOOGLE_SLIDES_MIME_TYPE,
                    },
                    media_body=media,
                    supportsAllDrives=True,
                    fields="id,name,mimeType,webViewLink",
                )
                .execute()
            )

        presentation_id = imported_file.get("id")
        if not presentation_id:
            raise ValueError("Google Drive did not return the imported file id")

        presentation = (
            get_slides_service()
            .presentations()
            .get(presentationId=presentation_id)
            .execute()
        )
        link = imported_file.get("webViewLink") or (
            f"https://docs.google.com/presentation/d/{presentation_id}/edit"
        )
        slides = presentation.get("slides", [])

        def validation_failed(message: str, **details: Any) -> str:
            cleanup_succeeded = _delete_imported_presentation(
                drive_service, presentation_id
            )
            payload: dict[str, Any] = {
                "status": "validation_failed",
                "message": message,
                "title": presentation.get("title") or resolved_title,
                **details,
            }
            if not cleanup_succeeded:
                payload["presentation_id"] = presentation_id
                payload["link"] = link
                payload["cleanup_failed"] = True
            return json.dumps(payload, ensure_ascii=False)

        if not slides:
            return validation_failed(
                "Imported presentation contains no slides",
                slide_count=0,
                empty_slide_numbers=[],
                missing_text_slide_numbers=[],
            )
        if len(expected_slide_text) != len(slides):
            return validation_failed(
                "Imported presentation slide count differs from the PPTX",
                slide_count=len(slides),
                expected_slide_count=len(expected_slide_text),
                empty_slide_numbers=[],
                missing_text_slide_numbers=[],
            )
        empty_slide_numbers = [
            index + 1
            for index, slide in enumerate(slides)
            if expected_slide_has_content[index] and _slide_is_empty(slide)
        ]
        actual_slide_text = [
            _normalize_comparison_text(" ".join(_slide_summary(slide, index)["text"]))
            for index, slide in enumerate(slides)
        ]
        missing_text_slide_numbers = [
            index + 1
            for index, expected_text in enumerate(expected_slide_text)
            if expected_text
            and not _contains_text_tokens(expected_text, actual_slide_text[index])
        ]
        if empty_slide_numbers or missing_text_slide_numbers:
            return validation_failed(
                "Imported presentation contains empty slides or missing text",
                slide_count=len(slides),
                empty_slide_numbers=empty_slide_numbers,
                missing_text_slide_numbers=missing_text_slide_numbers,
            )

        return json.dumps(
            {
                "status": "success",
                "presentation_id": presentation_id,
                "title": presentation.get("title") or resolved_title,
                "link": link,
                "slide_count": len(slides),
                "empty_slide_numbers": [],
                "missing_text_slide_numbers": [],
            },
            ensure_ascii=False,
        )
    except PermissionError as e:
        if drive_service is not None and presentation_id:
            _delete_imported_presentation(drive_service, presentation_id)
        logger.error("PPTX import rejected: %s", e)
        return _error(str(e))
    except FileNotFoundError as e:
        if drive_service is not None and presentation_id:
            _delete_imported_presentation(drive_service, presentation_id)
        logger.error("PPTX file not found: %s", e)
        return _error("PPTX file was not found")
    except ValueError as e:
        if drive_service is not None and presentation_id:
            _delete_imported_presentation(drive_service, presentation_id)
        logger.error("PPTX import rejected: %s", e)
        return _error(str(e))
    except OSError as e:
        if drive_service is not None and presentation_id:
            _delete_imported_presentation(drive_service, presentation_id)
        logger.error("Error reading PPTX during Google Slides import: %s", e)
        return _error("Could not read the PPTX file")
    except Exception as e:
        if drive_service is not None and presentation_id:
            _delete_imported_presentation(drive_service, presentation_id)
        logger.error(f"Error importing PPTX into Google Slides: {e}")
        return _error(str(e))


@mcp.tool()
def google_slides_add_slide(
    presentation_id: str,
    title: str = "",
    body: str = "",
    layout: _Layout = "TITLE_AND_BODY",
    default_slide_id: str = "",
    preserve_blank_slide: bool = False,
) -> str:
    """
    Append a slide with a title and body text to a Google Slides presentation.
    This is a text/layout primitive and does not create visual styling. For a
    polished or designed deck, generate a local PPTX with the presentation
    generator and call google_slides_import_pptx instead.
    The body supports plain text; use newlines to separate bullet lines — each
    line is rendered as its own bulleted paragraph (don't type literal "•"/"-"
    markers, Slides adds the bullet glyph itself — only "•"/"-"/"*" are
    recognized as markers to strip; other conventions like numbered lists
    ("1.") or other dash/arrow glyphs ("–"/"—"/"→") are inserted as
    literal text and will visually double up with Slides' own bullet).

    layout picks the slide's predefined layout, which decides which of
    title/body actually have a placeholder to land in:
      - "TITLE_AND_BODY" (default): title + full bulleted body content.
        body is required here — this layout rejects an empty body rather
        than silently creating a title-only slide. An empty title is
        still allowed as long as body has content (e.g. a title-less
        continuation slide) — this is intentional, not a gap.
      - "TITLE": a cover/section-opening slide — title is the big centered
        title, body (optional) becomes the subtitle line (not bulleted).
      - "TITLE_ONLY", "SECTION_HEADER": title only, no body placeholder —
        pass body="" or the call is rejected.
    For blank/custom slides, use google_slides_batch_update directly; this
    helper only exposes layouts with title/body placeholders.

    Layouts with a title placeholder that don't already require a body
    (TITLE, TITLE_ONLY, SECTION_HEADER — not TITLE_AND_BODY, whose own
    body-required rule above already covers it) still need at least a
    non-empty title — or, for
    "TITLE", a non-empty body/subtitle instead — since otherwise the call
    would produce a completely empty slide; that combination is rejected.

    This assumes the presentation uses a standard/default theme. A custom
    theme whose master doesn't expose the placeholder types listed above
    for a given layout will fail with a raw Slides API error from the
    underlying batchUpdate call rather than a friendly one — use
    google_slides_batch_update directly for presentations with a
    non-standard theme.

    When google_slides_create_presentation returned a default_slide_id, pass
    it on the first call: each call may run in a separate MCP process, and
    with the id the default page is removed in the same batch. An empty
    default page is removed even if other pages were added before this call;
    one that only shows the deck title that create wrote into it is removed
    only while it is the only page, and is kept as the cover once other pages
    exist. Without an id, only a sole page with no elements at all is
    removed, so a default page with title/subtitle placeholders stays as the
    first slide (showing the deck title when create wrote it); an empty page
    in a multi-page deck is preserved.
    Set preserve_blank_slide=True when that sole empty page is intentional.
    Unless the user specified a different first slide, make the first slide
    of a new deck its cover: layout="TITLE" with the deck title as title and
    an optional subtitle as body.

    To fix a slide this call already created (wrong/missing text), use
    google_slides_update_slide with its slide_id — do NOT call
    google_slides_add_slide again, that creates a second, duplicate slide
    rather than editing the first one.

    After adding all the slides for a deck, call
    google_slides_get_presentation once to read back every slide's actual
    title/body text and confirm it matches what you intended (e.g. against
    the user's outline) before telling the user the deck is done.
    """
    opened = False
    try:
        if not isinstance(layout, str):
            return _error(f"'layout' must be a string, got {type(layout).__name__}.")

        title = _normalize_title(title)
        body = _normalize_newlines(body)
        normalized_layout = layout.strip().upper()

        if normalized_layout not in _LAYOUT_PLACEHOLDERS:
            return _error(
                f"Unknown layout '{layout}'. Supported layouts: "
                f"{', '.join(sorted(_LAYOUT_PLACEHOLDERS))}"
            )

        title_placeholder, body_placeholder = _LAYOUT_PLACEHOLDERS[normalized_layout]
        is_bulleted = body_required = body_placeholder == "BODY"
        body = _prepare_body_text(body, is_bulleted)

        if title.strip() and title_placeholder is None:
            return _error(
                f"layout '{normalized_layout}' has no title placeholder, "
                "so 'title' would be silently dropped. Use a different "
                "layout, or google_slides_batch_update for a custom "
                "text box."
            )
        if body.strip() and body_placeholder is None:
            return _error(
                f"layout '{normalized_layout}' has no body placeholder, "
                "so 'body' would be silently dropped. Use a layout "
                "with a body/subtitle placeholder (e.g. "
                "TITLE_AND_BODY, TITLE) or omit body."
            )
        if body_required and not body.strip():
            return _error(
                f"layout '{normalized_layout}' expects body content but "
                "none was provided. Include this slide's full bullet/"
                "detail text in 'body' — don't create the slide "
                "with just a title."
            )
        if (
            title_placeholder is not None
            and not body_required
            and not title.strip()
            and not body.strip()
        ):
            return _error(
                f"layout '{normalized_layout}' needs at least a non-empty "
                "'title'"
                + (" or 'body'" if body_placeholder is not None else "")
                + " — as given, this call would create a completely "
                "empty slide."
            )

        pres_id = _resolve_presentation_id(presentation_id, editing=True)
        service = get_slides_service()

        tracked_default_slide_id = _CREATED_DEFAULT_SLIDES.get(pres_id)
        requested_default_slide_id = default_slide_id.strip()
        candidate_default_slide_id = (
            requested_default_slide_id or tracked_default_slide_id
        )
        if preserve_blank_slide:
            # Preserving the page consumes any in-process default inference;
            # otherwise a later call without this flag could delete the page
            # the caller explicitly asked us to keep.
            if candidate_default_slide_id:
                _PRESERVED_BLANK_SLIDES.setdefault(pres_id, set()).add(
                    candidate_default_slide_id
                )
            _CREATED_DEFAULT_SLIDES.pop(pres_id, None)
            candidate_default_slide_id = ""
        # A Slides API create call starts with a default blank page. Read the
        # current pages before creating the next real page, then delete that
        # default only after the new page has been created in this same atomic
        # batch. A known id can identify the default even after other pages
        # were added; without one, only a sole empty page is considered the
        # default, so user content in a multi-page deck is never guessed away.
        existing: dict[str, Any] = {"slides": []}
        if candidate_default_slide_id or not preserve_blank_slide:
            existing = service.presentations().get(presentationId=pres_id).execute()
            opened = True
        slide_to_remove = None
        existing_slides = existing.get("slides", [])
        if not preserve_blank_slide:
            if candidate_default_slide_id:
                candidate_slide = next(
                    (
                        slide
                        for slide in existing_slides
                        if slide.get("objectId") == candidate_default_slide_id
                    ),
                    None,
                )
                if candidate_slide is not None:
                    # A default page that create titled is already a usable
                    # cover. Replace it only on a call that passes its id while
                    # it is still the only page, as the descriptions ask for
                    # the first call. With only the id remembered in this
                    # process, or once other pages exist, it stays the cover.
                    if _slide_is_empty(candidate_slide) or (
                        requested_default_slide_id
                        and len(existing_slides) == 1
                        and _slide_is_titled_default(
                            candidate_slide, existing.get("title")
                        )
                    ):
                        slide_to_remove = candidate_default_slide_id
                    else:
                        # The caller used the page before asking us to append
                        # a slide, so it is no longer the untouched default.
                        _CREATED_DEFAULT_SLIDES.pop(pres_id, None)
            elif (
                len(existing_slides) == 1
                and existing_slides[0].get("objectId")
                not in _PRESERVED_BLANK_SLIDES.get(pres_id, set())
                and _slide_is_untouched_default(existing_slides[0])
            ):
                # Without a cross-process id, the only safe inference is a
                # presentation whose sole page is empty. Never infer that a
                # blank page in a multi-page deck is Google's default.
                slide_to_remove = existing_slides[0].get("objectId")

        if slide_to_remove:
            _PRESERVED_BLANK_SLIDES.get(pres_id, set()).discard(slide_to_remove)

        slide_id = f"slide_{uuid.uuid4().hex[:12]}"
        title_id = f"{slide_id}_title"
        body_id = f"{slide_id}_body"

        placeholder_mappings: list[dict[str, Any]] = []
        if title_placeholder is not None:
            placeholder_mappings.append(
                {
                    "layoutPlaceholder": {"type": title_placeholder},
                    "objectId": title_id,
                }
            )
        if body_placeholder is not None:
            placeholder_mappings.append(
                {
                    "layoutPlaceholder": {"type": body_placeholder},
                    "objectId": body_id,
                }
            )

        create_slide: dict[str, Any] = {
            "objectId": slide_id,
            "slideLayoutReference": {"predefinedLayout": normalized_layout},
        }
        if placeholder_mappings:
            create_slide["placeholderIdMappings"] = placeholder_mappings

        requests: list[dict[str, Any]] = [{"createSlide": create_slide}]
        if title.strip():
            requests.append(
                {"insertText": {"objectId": title_id, "text": title.strip()}}
            )
        if body.strip():
            requests.extend(_body_insert_requests(body_id, body, is_bulleted))
        if slide_to_remove:
            requests.append({"deleteObject": {"objectId": slide_to_remove}})

        service.presentations().batchUpdate(
            presentationId=pres_id, body={"requests": requests}
        ).execute()
        if slide_to_remove and _CREATED_DEFAULT_SLIDES.get(pres_id) == slide_to_remove:
            _CREATED_DEFAULT_SLIDES.pop(pres_id, None)

        return json.dumps(
            {
                "status": "success",
                "presentation_id": pres_id,
                "slide_id": slide_id,
                "layout": normalized_layout,
                "default_slide_removed": bool(slide_to_remove),
            },
            ensure_ascii=False,
        )
    except Exception as e:
        logger.error(f"Error adding slide: {e}")
        return _presentation_error(e, presentation_id, editing=True, opened=opened)


@mcp.tool()
def google_slides_update_slide(
    presentation_id: str, slide_id: str, title: str = "", body: str = ""
) -> str:
    """
    Replace the title and/or body text of an existing slide, identified by
    its slide_id — either the one a prior google_slides_add_slide call
    returned, or the "object_id" field google_slides_get_presentation
    reports for that slide (the same id under a different key) — instead
    of creating a new one. Use this to fix a slide that came out wrong or
    incomplete — calling google_slides_add_slide again does NOT edit that
    slide, it creates a duplicate one next to it.

    Only the placeholders you pass non-empty text for are touched; omit
    title or body to leave that placeholder untouched. Body text follows
    the same rules as google_slides_add_slide: newline-separated lines
    become bulleted paragraphs (with literal "•"/"-"/"*" markers and blank
    lines stripped) when the slide's body placeholder is a real "BODY"
    type, not a plain "SUBTITLE".
    """
    opened = False
    try:
        if not title and not body:
            return _error("Provide at least one of 'title' or 'body' to update.")

        title = _normalize_title(title)
        body = _normalize_newlines(body)

        if title and not title.strip():
            return _error("'title' is whitespace-only; provide real text or omit it.")
        if body and not body.strip():
            return _error("'body' is whitespace-only; provide real text or omit it.")

        pres_id = _resolve_presentation_id(presentation_id, editing=True)
        service = get_slides_service()
        presentation = service.presentations().get(presentationId=pres_id).execute()
        opened = True

        slide = _find_slide(presentation, slide_id)
        if slide is None:
            return _error(f"No slide with id '{slide_id}' in this presentation.")

        placeholders = _find_placeholders(slide)

        if title and "title" not in placeholders:
            return _error(
                f"Slide '{slide_id}' has no title placeholder, so 'title' has nowhere to go."
            )
        if body and "body" not in placeholders:
            return _error(
                f"Slide '{slide_id}' has no body/subtitle placeholder, so 'body' has nowhere to go."
            )

        if body:
            _, body_type = placeholders["body"]
            body = _prepare_body_text(body, body_type == "BODY")
            if not body:
                return _error(
                    "'body' has no real content once bullet markers/blank "
                    "lines are removed; provide actual detail text or "
                    "omit it."
                )

        requests: list[dict[str, Any]] = []
        for role, text in (("title", title.strip()), ("body", body)):
            if not text:
                continue
            element, placeholder_type = placeholders[role]
            object_id = element["objectId"]
            # Slides represents a "cleared" placeholder as a lone trailing
            # "\n" (the implicit paragraph terminator every text-containing
            # shape carries), not a truly empty string — treat exactly
            # that (or genuinely empty) as nothing-to-delete. Anything
            # else, even whitespace-only text beyond that terminator (e.g.
            # a stray " \n" from a slide edited by something other than
            # this tool), still gets deleteText: insertText has no
            # insertionIndex here, so it defaults to prepending at the
            # start rather than replacing — skipping the delete for
            # arbitrary whitespace would leave that stale text merged
            # into what's supposed to be a clean replacement.
            if _element_text(element) not in ("", "\n"):
                requests.append(
                    {
                        "deleteText": {
                            "objectId": object_id,
                            "textRange": {"type": "ALL"},
                        }
                    }
                )
            if role == "title":
                requests.append({"insertText": {"objectId": object_id, "text": text}})
            else:
                requests.extend(
                    _body_insert_requests(object_id, text, placeholder_type == "BODY")
                )

        service.presentations().batchUpdate(
            presentationId=pres_id, body={"requests": requests}
        ).execute()

        return json.dumps(
            {
                "status": "success",
                "presentation_id": pres_id,
                "slide_id": slide_id,
            },
            ensure_ascii=False,
        )
    except Exception as e:
        logger.error(f"Error updating slide: {e}")
        return _presentation_error(e, presentation_id, editing=True, opened=opened)


@mcp.tool()
def google_slides_delete_slide(presentation_id: str, slide_id: str) -> str:
    """
    Permanently delete one slide from a presentation by its slide_id.
    Use this to remove a duplicate or wrong slide — for example one created
    by calling google_slides_add_slide a second time instead of using
    google_slides_update_slide to fix the original.
    """
    opened = False
    try:
        pres_id = _resolve_presentation_id(presentation_id, editing=True)
        service = get_slides_service()
        presentation = service.presentations().get(presentationId=pres_id).execute()
        opened = True

        if _find_slide(presentation, slide_id) is None:
            return _error(
                f"No slide with id '{slide_id}' in this presentation — "
                "refusing to delete. Make sure this is a slide id, not a "
                "placeholder shape id (e.g. one ending in '_title'/'_body')."
            )

        service.presentations().batchUpdate(
            presentationId=pres_id,
            body={"requests": [{"deleteObject": {"objectId": slide_id}}]},
        ).execute()

        return json.dumps(
            {
                "status": "success",
                "presentation_id": pres_id,
                "slide_id": slide_id,
            },
            ensure_ascii=False,
        )
    except Exception as e:
        logger.error(f"Error deleting slide: {e}")
        return _presentation_error(e, presentation_id, editing=True, opened=opened)


@mcp.tool()
def google_slides_batch_update(presentation_id: str, requests_json: str) -> str:
    """
    Advanced: apply raw Google Slides API batchUpdate requests to a presentation.
    requests_json must be a JSON array of request objects following the Slides API
    schema (e.g. createShape, insertText, updateTextStyle, createImage).
    Use this only when the simpler tools cannot express the required change.
    When following google_slides_create_presentation across MCP processes, use
    the returned default_slide_id in a deleteObject request after creating the
    first content slide; this tool cannot infer that id from a later session.
    """
    try:
        pres_id = _resolve_presentation_id(presentation_id, editing=True)
        requests = json.loads(requests_json)
        if not isinstance(requests, list):
            raise ValueError("requests_json must be a JSON array of request objects")

        service = get_slides_service()
        result = (
            service.presentations()
            .batchUpdate(presentationId=pres_id, body={"requests": requests})
            .execute()
        )

        return json.dumps(
            {
                "status": "success",
                "presentation_id": result.get("presentationId", pres_id),
                "replies": result.get("replies", []),
            },
            ensure_ascii=False,
        )
    except Exception as e:
        logger.error(f"Error applying batch update: {e}")
        return _presentation_error(e, presentation_id, editing=True)


# Rendering one slide produces a short-lived signed image URL that is served
# by Google's user-content host, not by the Slides API.  Downloading it must
# therefore not reuse the OAuth transport, and the product bounds the transfer
# (10 MiB, 30 s, 3 redirects) so a malformed response cannot exhaust a task.
_THUMBNAIL_MAX_BYTES = 10 * 1024 * 1024
_THUMBNAIL_TIMEOUT_SECONDS = 30.0
_THUMBNAIL_MAX_REDIRECTS = 3
_THUMBNAIL_CHUNK_BYTES = 64 * 1024
# Google's getThumbnail documentation does not promise one fixed subdomain;
# the images are served from Google's user-content domain (lh3, lh7-rt, ...).
_THUMBNAIL_HOST_SUFFIX = "googleusercontent.com"
_THUMBNAIL_PNG_MIME_TYPE = "image/png"
_THUMBNAIL_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
_THUMBNAIL_SIZES = ("SMALL", "MEDIUM", "LARGE")
_THUMBNAIL_REQUEST_HEADERS = {
    "Accept": "image/png",
    "User-Agent": "Xagent-Google-Slides-Thumbnail/1.0",
}
_THUMBNAIL_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9_.-]")
_THUMBNAIL_FILENAME_STEM_LENGTH = 120
# The current task workspace, injected by the host exactly as for
# google_drive_download_file; a thumbnail is another generated deliverable
# under the same output directory.
_THUMBNAIL_OUTPUT_DIR_ENV_VAR = "XAGENT_GOOGLE_DRIVE_OUTPUT_DIR"
# The whole transfer (connect, headers, every redirect hop and the body) has
# one absolute, interruptible deadline.  The join below is the backstop for a
# read that a socket cannot unblock immediately, so it must stay short.
_THUMBNAIL_TIMEOUT_GRACE_SECONDS = 1.0
# Per socket operation caps; the absolute deadline above still shortens these.
_THUMBNAIL_SOCKET_TIMEOUT_SECONDS = 10.0
# A decoded thumbnail image is a small rendered page, not an arbitrary image
# file; the cap stops a crafted PNG header from making Pillow allocate hugely.
_THUMBNAIL_MAX_PIXEL_COUNT = 40_000_000
_THUMBNAIL_MAX_NAME_ATTEMPTS = 10_000

# The signed image URL is short-lived access material.  It must never be
# returned, logged (including by the transport's own debug logger) or embedded
# in an exception message, so every URL that reaches such a surface is stripped
# of its query/fragment first.
_URL_IN_TEXT = re.compile(r"https?://[^\s\"'<>]+")
_SECRET_TOKEN_PARAMETER = re.compile(
    r"(?i)\b(access_token|token|signature|sig|x-goog-[a-z0-9-]+)=([^&\s\"'<>]+)"
)


def _strip_url_query(url: str) -> str:
    try:
        parsed = urlsplit(url)
    except ValueError:
        return "<redacted URL>"
    if not parsed.query and not parsed.fragment:
        return url
    return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))


def _redact_thumbnail_secrets(text: object) -> str:
    """Remove signed URL query strings and token parameters from ``text``."""
    if not isinstance(text, str):
        return ""
    redacted = _URL_IN_TEXT.sub(lambda match: _strip_url_query(match.group(0)), text)
    return _SECRET_TOKEN_PARAMETER.sub(
        lambda match: f"{match.group(1)}=<redacted>", redacted
    )


def _redact_error_payload(payload: str) -> str:
    """Apply URL redaction to the message of a connector error payload."""
    try:
        parsed = json.loads(payload)
    except (TypeError, ValueError):
        return _error(_redact_thumbnail_secrets(payload))
    if not isinstance(parsed, dict):
        return _error(_redact_thumbnail_secrets(payload))
    message = parsed.get("message")
    if isinstance(message, str):
        parsed["message"] = _redact_thumbnail_secrets(message)
    return json.dumps(parsed, ensure_ascii=False)


class _ThumbnailSecretsFilter(logging.Filter):
    """Keep transport debug lines while removing the signed image URL query."""

    def filter(self, record: logging.LogRecord) -> bool:
        try:
            message = record.getMessage()
        except Exception:
            return True
        redacted = _redact_thumbnail_secrets(message)
        if redacted != message:
            record.msg = redacted
            record.args = ()
        return True


def _install_thumbnail_log_filter() -> None:
    log_filter = _ThumbnailSecretsFilter()
    for name in (
        "urllib3",
        "urllib3.connection",
        "urllib3.connectionpool",
        "urllib3.poolmanager",
        "requests",
    ):
        target = logging.getLogger(name)
        if not any(
            isinstance(existing, _ThumbnailSecretsFilter) for existing in target.filters
        ):
            target.addFilter(log_filter)


_install_thumbnail_log_filter()


class _ThumbnailCancelToken:
    """Lets the waiting caller interrupt a blocked transfer at its deadline."""

    def __init__(self) -> None:
        self._stopped = threading.Event()
        self._lock = threading.Lock()
        self._interrupter: Any = None

    @property
    def stopped(self) -> bool:
        return self._stopped.is_set()

    def set_interrupter(self, interrupter: Any) -> None:
        with self._lock:
            self._interrupter = interrupter
            stopped = self._stopped.is_set()
        if stopped:
            _run_interrupter(interrupter)

    def clear_interrupter(self) -> None:
        with self._lock:
            self._interrupter = None

    def stop(self) -> None:
        with self._lock:
            self._stopped.set()
            interrupter = self._interrupter
        if interrupter is not None:
            _run_interrupter(interrupter)


def _run_interrupter(interrupter: Any) -> None:
    try:
        interrupter()
    except Exception:
        pass


def _thumbnail_response_socket(response: Any) -> Any:
    """Return the socket a streaming response is reading from, if reachable."""
    raw = getattr(response, "raw", None)
    if raw is None:
        return None
    connection = getattr(raw, "_connection", None)
    sock = getattr(connection, "sock", None)
    if sock is not None:
        return sock
    file_object = getattr(getattr(getattr(raw, "_fp", None), "fp", None), "raw", None)
    return getattr(file_object, "_sock", None)


def _shutdown_thumbnail_response(response: Any) -> None:
    """Interrupt a read that is blocked on the response's socket."""
    sock = _thumbnail_response_socket(response)
    if sock is not None:
        try:
            sock.shutdown(socket.SHUT_RDWR)
        except OSError:
            pass
        return
    try:
        response.close()
    except Exception:
        pass


class _ThumbnailDownloadError(RuntimeError):
    """A failure while fetching a thumbnail image.

    The message is deliberately generic: the signed image URL must never reach
    the caller, a log line, or an exception's text.
    """


def _normalize_thumbnail_size(value: object) -> str:
    """Validate ``size`` inside FastMCP's schema layer as well.

    The function body repeats this check for direct Python callers, which
    bypass FastMCP validation.
    """
    if not isinstance(value, str):
        raise ValueError("'size' must be one of SMALL, MEDIUM, LARGE")
    normalized = value.strip().upper()
    if normalized not in _THUMBNAIL_SIZES:
        raise ValueError("'size' must be one of SMALL, MEDIUM, LARGE")
    return normalized


_ThumbnailSize = Annotated[
    Literal["SMALL", "MEDIUM", "LARGE"], BeforeValidator(_normalize_thumbnail_size)
]


def _strict_page_number(value: object) -> int | None:
    """Refuse every value that is not an actual integer.

    Pydantic's lax mode would otherwise accept a JSON ``true`` or ``"2"`` as a
    valid page and silently change which slide is returned.
    """
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError("'page_number' must be an integer")
    return value


_PageNumber = Annotated[int | None, BeforeValidator(_strict_page_number)]


def _is_allowed_thumbnail_host(hostname: str) -> bool:
    return hostname == _THUMBNAIL_HOST_SUFFIX or hostname.endswith(
        "." + _THUMBNAIL_HOST_SUFFIX
    )


def _validate_thumbnail_url(url: object) -> str:
    """Return ``url`` when it is an HTTPS Google user-content image URL.

    Every redirect hop is passed through this function, so a permitted first
    URL cannot hand the next hop to a plain client.
    """
    if not isinstance(url, str) or not url.strip():
        raise ValueError("thumbnail image URL is missing")
    if "\\" in url or any(
        character.isspace() or ord(character) < 32 for character in url
    ):
        raise ValueError("thumbnail image URL contains invalid characters")
    try:
        parsed = urlsplit(url)
    except ValueError:
        raise ValueError("thumbnail image URL could not be parsed") from None
    if parsed.scheme.lower() != "https":
        raise ValueError("thumbnail image URL must use HTTPS")
    if (
        parsed.username is not None
        or parsed.password is not None
        or "@" in parsed.netloc
    ):
        raise ValueError("thumbnail image URL must not include user information")
    hostname = parsed.hostname
    if not hostname:
        raise ValueError("thumbnail image URL has no host")
    try:
        port = parsed.port
    except ValueError:
        raise ValueError("thumbnail image URL has an invalid port") from None
    if port not in (None, 443):
        raise ValueError("thumbnail image URL must not use a non-standard port")
    host = hostname.rstrip(".").lower()
    try:
        ipaddress.ip_address(host)
    except ValueError:
        pass
    else:
        raise ValueError("thumbnail image URL host is not a Google image host")
    if host in {"localhost", "ip6-localhost"} or host.endswith(".localhost"):
        raise ValueError("thumbnail image URL host is not a Google image host")
    if not _is_allowed_thumbnail_host(host):
        raise ValueError("thumbnail image URL host is not a Google image host")
    return url


def _new_thumbnail_session() -> requests.Session:
    """A fresh public-only client for the image host.

    It carries no connector credentials, authorization or cookies, and it is
    separate from the Google API client's OAuth transport.
    """
    return requests.Session()


def _send_thumbnail_request(
    session: Any, prepared_request: Any, *, timeout: Any
) -> Any:
    """Send one handled hop without requests' redirect-body pre-read.

    ``requests.Session.send`` resolves ``allow_redirects=False`` by consuming
    the redirect response body to build ``Response.next``; that would buffer an
    unbounded body before the Location is validated.  Calling the mounted
    adapter directly keeps the streaming response unread and keeps the
    public-host checks of the shared client.
    """
    adapter = session.get_adapter(prepared_request.url)
    settings = session.merge_environment_settings(
        prepared_request.url, {}, True, None, None
    )
    return adapter.send(prepared_request, timeout=timeout, **settings)


def _thumbnail_hop_timeouts(deadline: float, cancel: _ThumbnailCancelToken) -> Any:
    if cancel.stopped:
        raise _ThumbnailDownloadError("thumbnail image download timed out")
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise _ThumbnailDownloadError("thumbnail image download timed out")
    operation_timeout = min(remaining, _THUMBNAIL_SOCKET_TIMEOUT_SECONDS)
    return operation_timeout, operation_timeout


def _validate_thumbnail_png(data: bytes) -> None:
    """Refuse a body that is only PNG-shaped unless it decodes as a PNG image.

    The URL can be served by an arbitrary host inside the allowed domain, so a
    signature prefix alone does not make the bytes a deliverable image.  Pillow
    is an existing dependency; decoding is bounded by a pixel-count cap first.
    """
    if not data.startswith(_THUMBNAIL_PNG_SIGNATURE):
        raise _ThumbnailDownloadError("downloaded thumbnail image is not a PNG image")
    try:
        from PIL import Image
    except Exception:
        raise _ThumbnailDownloadError(
            "downloaded thumbnail image is not a valid PNG image"
        ) from None
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.format != "PNG":
                raise _ThumbnailDownloadError(
                    "downloaded thumbnail image is not a PNG image"
                )
            width, height = image.size
            if width <= 0 or height <= 0 or width * height > _THUMBNAIL_MAX_PIXEL_COUNT:
                raise _ThumbnailDownloadError(
                    "downloaded thumbnail image has an unsupported size"
                )
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            image.load()
    except _ThumbnailDownloadError:
        raise
    except Exception:
        raise _ThumbnailDownloadError(
            "downloaded thumbnail image is not a valid PNG image"
        ) from None


def _read_thumbnail_body(
    response: Any, deadline: float, cancel: _ThumbnailCancelToken
) -> bytes:
    """Read a bounded body, then decode it as a PNG without trusting its name."""
    declared_length = response.headers.get("Content-Length")
    if declared_length is not None:
        trimmed_length = str(declared_length).strip()
        if not trimmed_length.isdigit():
            raise _ThumbnailDownloadError(
                "thumbnail image download returned an invalid Content-Length header"
            )
        if int(trimmed_length) > _THUMBNAIL_MAX_BYTES:
            raise _ThumbnailDownloadError("thumbnail image is too large to download")
    content_encoding = (
        str(response.headers.get("Content-Encoding", "identity")).strip().lower()
    )
    if content_encoding not in ("", "identity"):
        raise _ThumbnailDownloadError(
            "thumbnail image download used an unsupported content encoding"
        )
    chunks: list[bytes] = []
    total = 0
    try:
        for chunk in response.iter_content(_THUMBNAIL_CHUNK_BYTES):
            if cancel.stopped or time.monotonic() > deadline:
                raise _ThumbnailDownloadError("thumbnail image download timed out")
            if not chunk:
                continue
            total += len(chunk)
            if total > _THUMBNAIL_MAX_BYTES:
                raise _ThumbnailDownloadError(
                    "thumbnail image is too large to download"
                )
            chunks.append(chunk)
    except requests.RequestException as exc:
        if cancel.stopped or isinstance(exc, requests.Timeout):
            raise _ThumbnailDownloadError(
                "thumbnail image download timed out"
            ) from None
        raise _ThumbnailDownloadError(
            "thumbnail image download failed: network error"
        ) from None
    data = b"".join(chunks)
    if not data:
        raise _ThumbnailDownloadError("downloaded thumbnail image is empty")
    _validate_thumbnail_png(data)
    return data


def _download_thumbnail_worker(
    content_url: str,
    deadline: float,
    cancel: _ThumbnailCancelToken,
    result: dict[str, Any],
) -> None:
    """Download the image in a cancellable thread with a hard total deadline."""
    session = None
    try:
        session = _new_thumbnail_session()
        cancel.set_interrupter(session.close)
        current_url = content_url
        redirects_followed = 0
        while True:
            if cancel.stopped:
                raise _ThumbnailDownloadError("thumbnail image download timed out")
            hop_timeout = _thumbnail_hop_timeouts(deadline, cancel)
            prepared = session.prepare_request(
                requests.Request(
                    "GET", current_url, headers=dict(_THUMBNAIL_REQUEST_HEADERS)
                )
            )
            # No hop may carry cookies: the signed URL is the only credential.
            prepared.headers.pop("Cookie", None)
            try:
                response = _send_thumbnail_request(
                    session, prepared, timeout=hop_timeout
                )
            except requests.RequestException as exc:
                if cancel.stopped or isinstance(exc, requests.Timeout):
                    raise _ThumbnailDownloadError(
                        "thumbnail image download timed out"
                    ) from None
                raise _ThumbnailDownloadError(
                    "thumbnail image download failed: network error"
                ) from None
            cancel.set_interrupter(
                lambda response=response: _shutdown_thumbnail_response(response)
            )
            try:
                status = response.status_code
                if 300 <= status < 400:
                    location = response.headers.get("Location")
                    if not isinstance(location, str) or not location.strip():
                        raise _ThumbnailDownloadError(
                            "thumbnail image download returned an invalid redirect"
                        )
                    if redirects_followed >= _THUMBNAIL_MAX_REDIRECTS:
                        raise _ThumbnailDownloadError(
                            "thumbnail image download redirected too many times"
                        )
                    redirects_followed += 1
                    current_url = _validate_thumbnail_url(
                        urljoin(current_url, location.strip())
                    )
                    continue
                if not 200 <= status < 300:
                    raise _ThumbnailDownloadError(
                        f"thumbnail image download failed with HTTP {int(status)}"
                    )
                result["data"] = _read_thumbnail_body(response, deadline, cancel)
                return
            except requests.RequestException as exc:
                if cancel.stopped or isinstance(exc, requests.Timeout):
                    raise _ThumbnailDownloadError(
                        "thumbnail image download timed out"
                    ) from None
                raise _ThumbnailDownloadError(
                    "thumbnail image download failed: network error"
                ) from None
            finally:
                cancel.set_interrupter(session.close)
                response.close()
    except _ThumbnailDownloadError as error:
        result["error"] = error
    except ValueError as error:
        result["error"] = _ThumbnailDownloadError(str(error))
    except requests.RequestException:
        result["error"] = _ThumbnailDownloadError(
            "thumbnail image download timed out"
            if cancel.stopped
            else "thumbnail image download failed: network error"
        )
    except Exception:
        result["error"] = _ThumbnailDownloadError(
            "thumbnail image download timed out"
            if cancel.stopped
            else "thumbnail image download failed: network error"
        )
    finally:
        cancel.clear_interrupter()
        if session is not None:
            try:
                session.close()
            except Exception:
                pass


def _download_thumbnail_image(content_url: object) -> bytes:
    """Fetch and validate one thumbnail image under a hard total deadline."""
    validated_url = _validate_thumbnail_url(content_url)
    cancel = _ThumbnailCancelToken()
    deadline = time.monotonic() + _THUMBNAIL_TIMEOUT_SECONDS
    result: dict[str, Any] = {}
    worker = threading.Thread(
        target=_download_thumbnail_worker,
        args=(validated_url, deadline, cancel, result),
        name="google-slides-thumbnail",
        daemon=True,
    )
    worker.start()
    worker.join(_THUMBNAIL_TIMEOUT_SECONDS)
    if worker.is_alive():
        cancel.stop()
        worker.join(_THUMBNAIL_TIMEOUT_GRACE_SECONDS)
        raise _ThumbnailDownloadError("thumbnail image download timed out")
    if cancel.stopped:
        raise _ThumbnailDownloadError("thumbnail image download timed out")
    error = result.get("error")
    if isinstance(error, _ThumbnailDownloadError):
        raise error
    data = result.get("data")
    if not isinstance(data, bytes):
        raise _ThumbnailDownloadError("thumbnail image download failed: network error")
    return data


def _thumbnail_output_dir() -> Path:
    """Return the current task's output directory for the saved image."""
    base = os.environ.get(_THUMBNAIL_OUTPUT_DIR_ENV_VAR, "").strip()
    if not base:
        raise RuntimeError(
            "No task workspace configured for this connector "
            f"({_THUMBNAIL_OUTPUT_DIR_ENV_VAR} is unset) — "
            "google_slides_get_slide_thumbnail needs a task workspace to "
            "write into."
        )
    output_dir = Path(base).expanduser().resolve() / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def _thumbnail_filename(*, page_number: int | None, slide_object_id: str) -> str:
    if page_number is not None:
        stem = f"slide_{page_number}"
    else:
        sanitized = _THUMBNAIL_UNSAFE_FILENAME_CHARS.sub("_", slide_object_id).strip(
            "._"
        )
        stem = f"slide_{sanitized or 'image'}"
    return stem[:_THUMBNAIL_FILENAME_STEM_LENGTH] + ".png"


def _open_thumbnail_output_file(output_dir: Path, filename: str) -> tuple[Path, int]:
    """Create a new output file exclusively, never following a symlink.

    ``O_EXCL`` makes the name choice atomic under concurrent calls, and
    ``O_NOFOLLOW`` plus the parent check keep the descriptor inside the task
    output directory even when a same-named symlink already exists.
    """
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    resolved_output_dir = output_dir.resolve()
    for counter in range(_THUMBNAIL_MAX_NAME_ATTEMPTS):
        candidate_name = filename if counter == 0 else f"{stem} ({counter}){suffix}"
        candidate = output_dir / candidate_name
        try:
            descriptor = os.open(
                candidate,
                os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0),
                0o600,
            )
        except FileExistsError:
            continue
        except OSError:
            raise _ThumbnailDownloadError(
                "could not save the thumbnail image in the task workspace"
            ) from None
        if candidate.parent.resolve() != resolved_output_dir:
            os.close(descriptor)
            raise _ThumbnailDownloadError(
                "could not save the thumbnail image in the task workspace"
            )
        return candidate, descriptor
    raise _ThumbnailDownloadError("could not choose a unique thumbnail image file name")


def _write_thumbnail_png(data: bytes, *, output_dir: Path, filename: str) -> Path:
    output_path, descriptor = _open_thumbnail_output_file(output_dir, filename)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(data)
    except OSError:
        output_path.unlink(missing_ok=True)
        raise _ThumbnailDownloadError(
            "could not save the thumbnail image in the task workspace"
        ) from None
    return output_path


@mcp.tool()
def google_slides_get_slide_thumbnail(
    presentation_id: str,
    slide_id: str | None = None,
    page_number: _PageNumber = None,
    size: _ThumbnailSize = "LARGE",
) -> str:
    """
    Render one slide of a Google Slides presentation as a PNG image in the
    task workspace, so its look can be shown or inspected.

    This covers Google Slides presentations only; it does not render slides
    from an uploaded .pptx file.

    Provide exactly one of slide_id or page_number:
      - slide_id is the slide's object id, as returned by
        google_slides_get_presentation ("object_id") or by
        google_slides_add_slide ("slide_id"). It is always used as an object
        id: a numeric-looking value is never reinterpreted as a page number.
      - page_number is the slide's 1-based position in the presentation. The
        object id is looked up with presentations.get(fields="slides.objectId"),
        and a page outside the deck's slide count is rejected.

    size selects the image size Google renders: "LARGE" (the default),
    "MEDIUM", or "SMALL". The bytes are decoded as a PNG image before this tool
    reports success, so a body that only starts with the PNG signature is
    rejected instead of being saved as a deliverable.

    The result reports the saved PNG's path in the current task workspace
    output directory, its mime_type ("image/png"), its byte size, and the
    presentation and slide ids. Each call creates a new file atomically and
    never overwrites an existing output file. The short-lived thumbnail link
    Google returns is never included in the result, and the image download
    sends no cookie or Google credentials on any hop.

    The host attaches a durable file_ref (file_id) to a successful result
    after registering the saved file. Only a successful result with a
    trusted, non-empty file_ref can be shown or reused in later turns; if the
    result carries file_ref_error, report that the file is not durable instead
    of presenting it as delivered.

    Failures return an error result, never an empty success: an invalid,
    unknown or out-of-range slide, a presentation or API error, an image host
    or redirect that is not allowed, a timeout, an oversized download, or a
    response that is empty or not a valid PNG image.
    """
    opened = False
    try:
        if not isinstance(presentation_id, str) or not presentation_id.strip():
            return _error("'presentation_id' must be a non-empty string.")
        if slide_id is not None and not isinstance(slide_id, str):
            return _error("'slide_id' must be a string.")
        if page_number is not None and (
            isinstance(page_number, bool) or not isinstance(page_number, int)
        ):
            return _error("'page_number' must be an integer.")
        if not isinstance(size, str):
            return _error("'size' must be one of SMALL, MEDIUM, LARGE.")
        if slide_id is None and page_number is None:
            return _error(
                "Provide exactly one of slide_id (a slide object id) or "
                "page_number (1-based); neither was given."
            )
        if slide_id is not None and page_number is not None:
            return _error(
                "Provide exactly one of slide_id (a slide object id) or "
                "page_number (1-based); both were given."
            )
        normalized_size = size.strip().upper()
        if normalized_size not in _THUMBNAIL_SIZES:
            return _error(
                f"Unknown size '{size}'. Supported sizes: SMALL, MEDIUM, LARGE."
            )
        if slide_id is not None:
            slide_id = slide_id.strip()
            if not slide_id:
                return _error("'slide_id' must not be empty.")
        if page_number is not None and page_number < 1:
            return _error("'page_number' must be 1 or greater.")

        output_dir = _thumbnail_output_dir()
        pres_id = _resolve_presentation_id(presentation_id)
        service = get_slides_service()

        target_slide_id: str
        if page_number is not None:
            presentation = (
                service.presentations()
                .get(presentationId=pres_id, fields="slides.objectId")
                .execute()
            )
            opened = True
            slides = (
                presentation.get("slides", []) if isinstance(presentation, dict) else []
            )
            object_ids: list[str] = [
                str(slide.get("objectId"))
                for slide in slides
                if isinstance(slide, dict)
                and isinstance(slide.get("objectId"), str)
                and slide.get("objectId")
            ]
            if not object_ids:
                return _error("The presentation has no slides to render.")
            if page_number > len(object_ids):
                return _error(
                    f"page_number {page_number} is out of range: the "
                    f"presentation has {len(object_ids)} slides."
                )
            target_slide_id = object_ids[page_number - 1]
        elif slide_id is not None:
            target_slide_id = slide_id
        else:
            return _error(
                "Provide exactly one of slide_id (a slide object id) or "
                "page_number (1-based); neither was given."
            )

        thumbnail = (
            service.presentations()
            .pages()
            .getThumbnail(
                presentationId=pres_id,
                pageObjectId=target_slide_id,
                thumbnailProperties_mimeType="PNG",
                thumbnailProperties_thumbnailSize=normalized_size,
            )
            .execute()
        )
        content_url = (
            thumbnail.get("contentUrl") if isinstance(thumbnail, dict) else None
        )
        if not isinstance(content_url, str) or not content_url.strip():
            return _error(
                "Google Slides did not return a thumbnail image URL for this slide."
            )

        data = _download_thumbnail_image(content_url)
        output_path = _write_thumbnail_png(
            data,
            output_dir=output_dir,
            filename=_thumbnail_filename(
                page_number=page_number, slide_object_id=target_slide_id
            ),
        )
        width = thumbnail.get("width")
        height = thumbnail.get("height")
        return json.dumps(
            {
                "status": "success",
                "presentation_id": pres_id,
                "slide_id": target_slide_id,
                "page_number": page_number,
                "mime_type": _THUMBNAIL_PNG_MIME_TYPE,
                "size": len(data),
                "path": str(output_path),
                "width": width if isinstance(width, int) else None,
                "height": height if isinstance(height, int) else None,
            },
            ensure_ascii=False,
        )
    except (ValueError, _ThumbnailDownloadError) as e:
        logger.error("Could not render slide thumbnail: %s", type(e).__name__)
        return _error(_redact_thumbnail_secrets(str(e)))
    except Exception as e:
        logger.error("Error rendering slide thumbnail: %s", type(e).__name__)
        return _redact_error_payload(
            _presentation_error(e, presentation_id, opened=opened)
        )


if __name__ == "__main__":
    mcp.run()
