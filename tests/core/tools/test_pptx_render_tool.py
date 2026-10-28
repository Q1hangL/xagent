"""Tests for rendering a workspace .pptx into one PNG per slide.

The conversion subprocess and (for most cases) the PDF raster step are mocked
so the selection, pixel budget, registration and cleanup contracts can be
checked deterministically. ``test_render_two_page_deck_with_real_soffice``
covers the real LibreOffice + PyMuPDF path and only skips when LibreOffice is
genuinely absent from this host's PATH.
"""

import asyncio
import base64
import hashlib
import math
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile
import threading
from types import SimpleNamespace

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import xagent.core.storage.manager as storage_manager
import xagent.web.models.database as database_module
from xagent.core.file_storage.factory import get_unscoped_file_storage
from xagent.core.workspace import TaskWorkspace
from xagent.web.models import Base
from xagent.web.models.task import Task
from xagent.web.models.uploaded_file import UploadedFile
from xagent.web.models.user import User

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
    "/x8AAwMCAO+aXioAAAAASUVORK5CYII="
)
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _load_render_module():
    from xagent.core.tools.core import pptx_render

    return pptx_render


def _load_conversion_module():
    from xagent.core.tools.core import pptx_conversion

    return pptx_conversion


@pytest.fixture
def durable_workspace(monkeypatch, tmp_path):
    """A real task-owned workspace backed by an isolated sqlite database."""
    uploads = tmp_path / "uploads"
    objects = tmp_path / "objects"
    storage_root = tmp_path / "storage"
    monkeypatch.setenv("XAGENT_UPLOADS_DIR", str(uploads))
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", objects.as_uri())
    monkeypatch.setenv("XAGENT_STORAGE_ROOT", str(storage_root))
    get_unscoped_file_storage.cache_clear()

    engine = create_engine(
        f"sqlite:///{tmp_path / 'metadata.db'}",
        connect_args={"check_same_thread": False},
    )
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    with session_factory() as db:
        user = User(username="render-owner", password_hash="hash")
        db.add(user)
        db.flush()
        db.add(Task(id=4242, user_id=user.id, title="render test task"))
        db.commit()
        user_pk = int(user.id)
    monkeypatch.setattr(database_module, "_SessionLocal", session_factory)
    monkeypatch.setattr(storage_manager, "create_db_session", session_factory)
    workspace = TaskWorkspace(
        id="web_task_4242",
        base_dir=str(uploads / f"user_{user_pk}"),
    )
    yield SimpleNamespace(
        workspace=workspace,
        session_factory=session_factory,
        base_dir=uploads / f"user_{user_pk}",
        storage_root=storage_root,
    )
    engine.dispose()
    get_unscoped_file_storage.cache_clear()


def _write_deck(directory: pathlib.Path, name: str = "deck.pptx") -> pathlib.Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    path.write_bytes(b"PK\x03\x04 synthetic presentation bytes")
    return path


def _arg_value(arguments, flag):
    return arguments[arguments.index(flag) + 1]


class _FakeMatrix:
    def __init__(self, x: float, y: float) -> None:
        self.x = x
        self.y = y
        self.a = x
        self.b = y

    def __eq__(self, other: object) -> bool:
        return isinstance(other, _FakeMatrix) and (self.x, self.y) == (other.x, other.y)


class _FakePixelExtent:
    def __init__(self, width: int, height: int) -> None:
        self.width = width
        self.height = height

    def save(self, path, output=None) -> None:
        if hasattr(path, "write"):
            path.write(PNG_1X1)
        else:
            pathlib.Path(path).write_bytes(PNG_1X1)

    def pil_save(self, stream, **kwargs) -> None:
        self.save(stream)


class _FakeRect:
    """Rect double that mirrors PyMuPDF's matrix/irect round trip."""

    def __init__(self, width: float, height: float):
        self.width = width
        self.height = height

    def __mul__(self, matrix):
        width = max(1, math.ceil(self.width * matrix.a - 1e-9))
        height = max(1, math.ceil(self.height * matrix.b - 1e-9))
        return SimpleNamespace(irect=SimpleNamespace(width=width, height=height))

    __rmul__ = __mul__


class _FakePage:
    def __init__(self, width: float, height: float, events: list, number: int):
        self.rect = _FakeRect(width, height)
        self._events = events
        self._number = number

    def get_pixmap(self, *, matrix, alpha):
        self._events.append(("get_pixmap", self._number, matrix))
        width = max(1, math.ceil(self.rect.width * matrix.a - 1e-9))
        height = max(1, math.ceil(self.rect.height * matrix.b - 1e-9))
        return _FakePixelExtent(width, height)


class _FakeDocument:
    def __init__(self, pages: list[_FakePage]):
        self._pages = pages
        self.page_count = len(pages)
        self.closed = False

    def __getitem__(self, index: int) -> _FakePage:
        return self._pages[index]

    def close(self) -> None:
        self.closed = True


class _FakeFitz:
    def __init__(self, pages: list[tuple[float, float]], events: list):
        self.pages = [
            _FakePage(width, height, events, number)
            for number, (width, height) in enumerate(pages, start=1)
        ]
        self.events = events
        self.documents: list[_FakeDocument] = []

    def open(self, _path):
        document = _FakeDocument(self.pages)
        self.documents.append(document)
        return document

    @staticmethod
    def Matrix(x: float, y: float) -> _FakeMatrix:  # noqa: N802 - mirrors PyMuPDF
        return _FakeMatrix(x, y)


def _install_conversion(monkeypatch, render, *, calls, timeout_records):
    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        calls.append(pathlib.Path(pptx_path))
        timeout_records.append(timeout_seconds)
        pdf_path = output_dir / (pathlib.Path(pptx_path).stem + ".pdf")
        pdf_path.write_bytes(b"%PDF-1.4 synthetic")
        return pdf_path

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)
    return calls


def _install_rasterizer(monkeypatch, render, *, page_count, requested, rendered):
    from xagent.core.tools.core.pptx_render import RenderedPage

    def render_pdf_pages_to_pngs(
        pdf_path,
        page_numbers,
        output_dir,
        *,
        output_fd,
        cancelled,
    ):
        requested.append(list(page_numbers))
        output_dir = pathlib.Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        pages = []
        for number in page_numbers:
            path = output_dir / f"slide_{number:03d}.png"
            path.write_bytes(PNG_1X1)
            pages.append(RenderedPage(page_number=number, path=path, width=1, height=1))
        rendered.append(pages)
        return pages

    monkeypatch.setattr(render, "render_pdf_pages_to_pngs", render_pdf_pages_to_pngs)
    monkeypatch.setattr(
        render, "_open_pdf", lambda _path: _FakeDocument([None] * page_count)
    )


def _output_rows(session_factory):
    with session_factory() as db:
        return list(db.query(UploadedFile).order_by(UploadedFile.filename).all())


def _register_input_deck(workspace, deck_path):
    return workspace.register_files([(str(deck_path), None)])[0]


def test_resolve_page_selection_defaults_to_the_first_twenty():
    render = _load_render_module()

    selection = render.resolve_page_selection(25, None, None)

    assert list(selection.page_numbers) == list(range(1, 21))
    assert selection.truncated is True
    assert selection.total_pages == 25

    small = render.resolve_page_selection(2, None, None)
    assert list(small.page_numbers) == [1, 2]
    assert small.truncated is False


def test_resolve_page_selection_honours_explicit_ranges():
    render = _load_render_module()

    assert list(render.resolve_page_selection(21, 1, 1).page_numbers) == [1]
    assert list(render.resolve_page_selection(21, 3, 5).page_numbers) == [3, 4, 5]
    assert list(render.resolve_page_selection(21, 21, 21).page_numbers) == [21]
    assert list(render.resolve_page_selection(21, 20, None).page_numbers) == [
        20,
        21,
    ]
    assert list(render.resolve_page_selection(21, None, 2).page_numbers) == [1, 2]


def test_resolve_page_selection_rejects_out_of_bounds_ranges():
    render = _load_render_module()

    with pytest.raises(render.PptxRenderError):
        render.resolve_page_selection(5, 4, 6)
    with pytest.raises(render.PptxRenderError):
        render.resolve_page_selection(5, 6, 6)
    with pytest.raises(render.PptxRenderError):
        render.resolve_page_selection(0, None, None)


def test_unrendered_page_ranges_report_skipped_pages():
    render = _load_render_module()

    assert render.unrendered_page_ranges(25, list(range(1, 21))) == [[21, 25]]
    assert render.unrendered_page_ranges(10, [3, 4, 5]) == [[1, 2], [6, 10]]
    assert render.unrendered_page_ranges(5, [1, 2, 3, 4, 5]) == []


@pytest.mark.asyncio
async def test_render_default_renders_every_page_of_a_small_deck(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    conversion_calls: list[pathlib.Path] = []
    timeouts: list[float] = []
    requested: list[list[int]] = []
    rendered: list[list] = []
    _install_conversion(
        monkeypatch,
        render,
        calls=conversion_calls,
        timeout_records=timeouts,
    )
    _install_rasterizer(
        monkeypatch, render, page_count=2, requested=requested, rendered=rendered
    )

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    assert result["total_pages"] == 2
    assert result["page_numbers"] == [1, 2]
    assert result["rendered_pages"] == 2
    assert result["truncated"] is False
    assert result["unrendered_pages"] == 0
    assert result["unrendered_page_ranges"] == []
    assert result["page_limit"] == render.MAX_RENDERED_PAGES
    assert result["max_pixels_per_page"] == render.MAX_PAGE_PIXELS
    assert result["conversion_timeout_seconds"] == 60
    assert requested == [[1, 2]]
    assert [page["page_number"] for page in result["pages"]] == [1, 2]
    for page in result["pages"]:
        assert page["file_ref"]["mime_type"] == "image/png"
        assert page["file_ref"]["filename"] == f"slide_{page['page_number']:03d}.png"
        assert page["file_ref"]["file_id"]
        assert page["artifacts"] == [
            {
                "type": "image",
                "file_id": page["file_ref"]["file_id"],
                "filename": page["file_ref"]["filename"],
                "mime_type": "image/png",
                "display": "inline",
            }
        ]
    assert len(result["artifacts"]) == 2
    assert timeouts == [60.0]


@pytest.mark.asyncio
async def test_render_default_caps_at_twenty_pages_without_rasterizing_more(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    requested: list[list[int]] = []
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(
        monkeypatch, render, page_count=25, requested=requested, rendered=[]
    )

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    assert requested == [list(range(1, 21))]
    assert result["page_limit"] == 20
    assert result["rendered_pages"] == 20
    assert result["truncated"] is True
    assert result["total_pages"] == 25
    assert result["unrendered_pages"] == 5
    assert result["unrendered_page_ranges"] == [[21, 25]]
    assert len(_output_rows(durable_workspace.session_factory)) == 20


@pytest.mark.asyncio
async def test_render_rejects_an_explicit_range_over_the_page_limit(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    conversions: list[pathlib.Path] = []
    _install_conversion(monkeypatch, render, calls=conversions, timeout_records=[])
    requested: list[list[int]] = []
    _install_rasterizer(
        monkeypatch, render, page_count=30, requested=requested, rendered=[]
    )

    result = await render.render_pptx_to_images(
        str(deck), durable_workspace.workspace, start_page=1, end_page=21
    )

    assert result["success"] is False
    assert result["error_code"] == "invalid_argument"
    assert "20" in result["error"]
    assert requested == []
    assert not list(durable_workspace.workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_honours_an_explicit_page_range(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    requested: list[list[int]] = []
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(
        monkeypatch, render, page_count=10, requested=requested, rendered=[]
    )

    result = await render.render_pptx_to_images(
        str(deck), durable_workspace.workspace, start_page=3, end_page=5
    )

    assert result["success"] is True
    assert requested == [[3, 4, 5]]
    assert result["page_numbers"] == [3, 4, 5]
    assert result["truncated"] is False
    assert result["unrendered_pages"] == 7
    assert result["unrendered_page_ranges"] == [[1, 2], [6, 10]]
    assert [page["file_ref"]["filename"] for page in result["pages"]] == [
        "slide_003.png",
        "slide_004.png",
        "slide_005.png",
    ]


@pytest.mark.asyncio
async def test_render_rejects_an_explicit_range_beyond_the_deck(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    requested: list[list[int]] = []
    _install_rasterizer(
        monkeypatch, render, page_count=5, requested=requested, rendered=[]
    )

    result = await render.render_pptx_to_images(
        str(deck), durable_workspace.workspace, start_page=4, end_page=6
    )

    assert result["success"] is False
    assert result["error_code"] == "page_range_invalid"
    assert "6" in result["error"]
    assert requested == []
    assert not list(durable_workspace.workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_reports_an_empty_presentation(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    requested: list[list[int]] = []
    _install_rasterizer(
        monkeypatch, render, page_count=0, requested=requested, rendered=[]
    )

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "empty_presentation"
    assert requested == []
    assert not list(durable_workspace.workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_page_limit_is_applied_before_any_pixmap_allocation(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    events: list = []
    fake_fitz = _FakeFitz([(720.0, 540.0)] * 25, events)
    monkeypatch.setitem(sys.modules, "fitz", fake_fitz)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    assert result["rendered_pages"] == 20
    pixmap_pages = [event[1] for event in events if event[0] == "get_pixmap"]
    assert pixmap_pages == list(range(1, 21))
    assert len(_output_rows(durable_workspace.session_factory)) == 20


@pytest.mark.asyncio
async def test_render_limits_pixels_before_requesting_a_huge_pixmap(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    events: list = []
    fake_fitz = _FakeFitz([(1_000_000.0, 1_000_000.0)], events)
    monkeypatch.setitem(sys.modules, "fitz", fake_fitz)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    pixmap_events = [event for event in events if event[0] == "get_pixmap"]
    assert len(pixmap_events) == 1
    assert len(fake_fitz.documents) == 1
    matrix = pixmap_events[0][2]
    assert matrix.a == matrix.b
    predicted_width = math.ceil(1_000_000.0 * matrix.a - 1e-9)
    predicted_height = math.ceil(1_000_000.0 * matrix.b - 1e-9)
    assert predicted_width * predicted_height <= render.MAX_PAGE_PIXELS
    assert matrix.a < 1.0


def _rasterize(render, pdf_path, numbers, output_dir):
    document = render._open_pdf(pdf_path)
    descriptor = os.open(output_dir, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        pages = render.render_pdf_pages_to_pngs(
            document,
            numbers,
            output_dir,
            output_fd=descriptor,
            cancelled=threading.Event(),
        )
        return document.page_count, pages
    finally:
        os.close(descriptor)
        document.close()


def test_real_rasterizer_stays_within_the_pixel_budget(tmp_path):
    fitz = pytest.importorskip("fitz")
    render = _load_render_module()
    document = fitz.open()
    try:
        page = document.new_page(width=720, height=540)
        page.insert_text((72, 72), "budget")
        pdf_path = tmp_path / "budget.pdf"
        document.save(str(pdf_path))
    finally:
        document.close()
    output_dir = tmp_path / "pngs"
    output_dir.mkdir()

    total_pages, pages = _rasterize(render, pdf_path, [1], output_dir)

    assert total_pages == 1
    assert len(pages) == 1
    rendered_page = pages[0]
    assert rendered_page.page_number == 1
    assert rendered_page.path.read_bytes().startswith(PNG_MAGIC)
    assert rendered_page.width * rendered_page.height <= render.MAX_PAGE_PIXELS
    assert rendered_page.width > 720
    assert abs(rendered_page.width / rendered_page.height - 720 / 540) < 0.01


def test_real_rasterizer_renders_only_the_selected_pages(tmp_path):
    fitz = pytest.importorskip("fitz")
    render = _load_render_module()
    document = fitz.open()
    try:
        for number in range(5):
            page = document.new_page(width=720, height=540)
            page.insert_text((72, 72), f"slide {number + 1}")
        pdf_path = tmp_path / "deck.pdf"
        document.save(str(pdf_path))
    finally:
        document.close()
    output_dir = tmp_path / "pngs"
    output_dir.mkdir()

    total_pages, pages = _rasterize(render, pdf_path, [2, 4], output_dir)

    assert total_pages == 5
    assert [page.page_number for page in pages] == [2, 4]
    assert sorted(path.name for path in output_dir.iterdir()) == [
        "slide_002.png",
        "slide_004.png",
    ]
    for page in pages:
        data = page.path.read_bytes()
        assert data.startswith(PNG_MAGIC)
        assert page.width > 0 and page.height > 0


@pytest.mark.asyncio
async def test_render_each_page_is_durably_registered_and_resolvable(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=2, requested=[], rendered=[])

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    rows = _output_rows(durable_workspace.session_factory)
    assert [row.filename for row in rows] == ["slide_001.png", "slide_002.png"]
    assert {row.mime_type for row in rows} == {"image/png"}
    assert {row.task_id for row in rows} == {4242}
    assert not any(str(row.filename).endswith(".pdf") for row in rows)
    fresh_workspace = TaskWorkspace(
        id="web_task_4242", base_dir=str(durable_workspace.base_dir)
    )
    for page in result["pages"]:
        file_id = page["file_ref"]["file_id"]
        resolved = fresh_workspace.resolve_file_id_detached(file_id)
        assert resolved is not None
        assert resolved.is_file()
        assert resolved.name == page["file_ref"]["filename"]
        assert page["file_ref"]["preview_url"] == f"/api/files/preview/{file_id}"
        assert page["file_ref"]["markdown_link"] is not None
        assert f"file:{file_id}" in page["file_ref"]["markdown_link"]
        assert page["artifacts"][0]["file_id"] == file_id


@pytest.mark.asyncio
async def test_render_requires_a_persistent_task_owner(monkeypatch, tmp_path):
    render = _load_render_module()
    deck = _write_deck(tmp_path)
    workspace = TaskWorkspace(id="scratch_workspace", base_dir=str(tmp_path / "ws"))
    conversions: list[pathlib.Path] = []
    _install_conversion(monkeypatch, render, calls=conversions, timeout_records=[])

    result = await render.render_pptx_to_images(str(deck), workspace)

    assert result["success"] is False
    assert result["error_code"] == "durable_registration_unavailable"
    assert "persistent" in result["error"] or "task owner" in result["error"]
    assert conversions == []
    assert not list(workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_requires_a_workspace(monkeypatch, tmp_path):
    render = _load_render_module()
    deck = _write_deck(tmp_path)

    result = await render.render_pptx_to_images(str(deck), None)

    assert result["success"] is False
    assert result["error_code"] == "workspace_unavailable"


@pytest.mark.asyncio
async def test_render_accepts_persistent_file_ids_and_allowed_paths(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    deck_bytes = deck.read_bytes()
    file_id = _register_input_deck(durable_workspace.workspace, deck)
    conversions: list[pathlib.Path] = []
    _install_conversion(monkeypatch, render, calls=conversions, timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])
    selectors = [
        f"file:{file_id}",
        file_id,
        "input/deck.pptx",
        str(deck),
    ]

    for selector in selectors:
        result = await render.render_pptx_to_images(
            selector, durable_workspace.workspace
        )
        assert result["success"] is True, (selector, result)
        assert pathlib.Path(result["pptx_path"]).resolve() == deck.resolve()

    assert len(conversions) == 1
    assert all(path != deck and not path.exists() for path in conversions)
    assert (
        hashlib.sha256(deck.read_bytes()).hexdigest()
        == hashlib.sha256(deck_bytes).hexdigest()
    )


@pytest.mark.asyncio
async def test_render_rejects_non_pptx_missing_outside_and_symlink_inputs(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    conversions: list[pathlib.Path] = []
    _install_conversion(monkeypatch, render, calls=conversions, timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])
    workspace = durable_workspace.workspace
    not_pptx = workspace.input_dir / "notes.txt"
    not_pptx.write_bytes(b"not a presentation")
    directory_input = workspace.input_dir / "folder.pptx"
    directory_input.mkdir()
    missing = workspace.input_dir / "missing.pptx"
    outside = tmp_path / "outside" / "outside.pptx"
    outside.parent.mkdir()
    outside.write_bytes(b"PK\x03\x04 outside")
    escape = workspace.input_dir / "escape.pptx"
    escape.symlink_to(outside)

    cases = [
        (str(not_pptx), "unsupported_format"),
        (str(directory_input), "input_not_a_file"),
        (str(missing), "input_not_found"),
        (str(outside), "input_outside_workspace"),
        (str(escape), "input_outside_workspace"),
        ("file:00000000-0000-0000-0000-000000000000", "input_not_found"),
    ]
    for selector, expected_code in cases:
        result = await render.render_pptx_to_images(selector, workspace)
        assert result["success"] is False, selector
        assert result["error_code"] == expected_code, (selector, result)

    assert conversions == []
    assert not list(workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_accepts_an_uppercase_pptx_suffix(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    _write_deck(durable_workspace.workspace.input_dir, "DECK.PPTX")
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])

    result = await render.render_pptx_to_images(
        "input/DECK.PPTX", durable_workspace.workspace
    )

    assert result["success"] is True


@pytest.mark.asyncio
async def test_render_passes_the_fixed_timeout_to_the_converter(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    conversion = _load_conversion_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    timeouts: list[float] = []
    _install_conversion(monkeypatch, render, calls=[], timeout_records=timeouts)
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    assert timeouts == [conversion.DEFAULT_PPTX_CONVERSION_TIMEOUT_SECONDS]


@pytest.mark.asyncio
async def test_render_cancellation_reclaims_the_subprocess_and_cleans_up(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    started = asyncio.Event()
    seen: dict[str, pathlib.Path] = {}

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        seen["temp_root"] = pathlib.Path(output_dir)
        seen["pdf_dir"] = pathlib.Path(output_dir)
        started.set()
        await asyncio.sleep(60)
        raise AssertionError("conversion must be cancelled")

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)

    task = asyncio.create_task(
        render.render_pptx_to_images(str(deck), durable_workspace.workspace)
    )
    await asyncio.wait_for(started.wait(), timeout=5)
    output_dirs = list(durable_workspace.workspace.output_dir.iterdir())
    assert output_dirs == []
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task

    assert not seen["pdf_dir"].exists()
    assert not seen["temp_root"].exists()
    assert _output_rows(durable_workspace.session_factory) == []


@pytest.mark.asyncio
@pytest.mark.parametrize("same_source", [False, True])
async def test_concurrent_renders_use_isolated_directories(
    durable_workspace, monkeypatch, tmp_path, same_source
):
    render = _load_render_module()
    first = _write_deck(durable_workspace.workspace.input_dir, "first.pptx")
    second = (
        first
        if same_source
        else _write_deck(durable_workspace.workspace.input_dir, "second.pptx")
    )
    seen: list[dict] = []
    release = asyncio.Event()
    arrived = asyncio.Event()
    in_flight = 0

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        nonlocal in_flight
        in_flight += 1
        if in_flight == 2:
            arrived.set()
        seen.append(
            {
                "source": pathlib.Path(pptx_path),
                "output_dir": pathlib.Path(output_dir),
                "temp_root": pathlib.Path(output_dir),
            }
        )
        await release.wait()
        pdf_path = pathlib.Path(output_dir) / (pathlib.Path(pptx_path).stem + ".pdf")
        pdf_path.write_bytes(b"%PDF-1.4 synthetic")
        return pdf_path

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])

    both = asyncio.gather(
        render.render_pptx_to_images(str(first), durable_workspace.workspace),
        render.render_pptx_to_images(str(second), durable_workspace.workspace),
    )
    await asyncio.wait_for(arrived.wait(), timeout=5)
    release.set()
    results = await asyncio.wait_for(both, timeout=10)

    assert [result["success"] for result in results] == [True, True]
    directories = [pathlib.Path(result["output_dir"]) for result in results]
    pdf_directories = [entry["output_dir"] for entry in seen]
    temp_roots = [entry["temp_root"] for entry in seen]
    assert len({str(path) for path in directories}) == 2
    assert len({str(path) for path in pdf_directories}) == 2
    assert len({str(path) for path in temp_roots}) == 2
    for path in directories:
        assert path.is_dir()
        assert (path / "slide_001.png").is_file()
    for path in temp_roots:
        assert not path.exists()
    assert len(_output_rows(durable_workspace.session_factory)) == 2
    cached = list(durable_workspace.storage_root.rglob("*.hidden.preview.pdf"))
    assert len(cached) == (1 if same_source else 2)
    assert all(path.read_bytes() == b"%PDF-1.4 synthetic" for path in cached)


@pytest.mark.asyncio
async def test_render_fails_clearly_when_libreoffice_is_missing(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    conversion = _load_conversion_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        raise conversion.PptxConverterUnavailableError(
            "LibreOffice is not installed or not on PATH"
        )

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "pdf_conversion_unavailable"
    assert "LibreOffice" in result["error"]
    assert not list(durable_workspace.workspace.output_dir.iterdir())
    assert _output_rows(durable_workspace.session_factory) == []


@pytest.mark.asyncio
async def test_render_reports_conversion_timeout_and_cleans_up(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    conversion = _load_conversion_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        raise conversion.PptxConversionTimeoutError(
            "LibreOffice conversion timed out after 60 seconds"
        )

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "pdf_conversion_timeout"
    assert not list(durable_workspace.workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_reports_conversion_failure_and_cleans_up(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    conversion = _load_conversion_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        raise conversion.PptxConversionFailedError("soffice exited 1")

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "pdf_conversion_failed"
    assert not list(durable_workspace.workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_render_reports_a_bad_pdf_and_cleans_up(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)

    async def convert_pptx_to_pdf(
        pptx_path, output_dir, *, timeout_seconds=60.0, include_hidden_slides=False
    ):
        pdf_path = pathlib.Path(output_dir) / (pathlib.Path(pptx_path).stem + ".pdf")
        pdf_path.write_bytes(b"this is not a pdf")
        return pdf_path

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert_pptx_to_pdf)

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] in {"pdf_unreadable", "raster_failed"}
    assert not list(durable_workspace.workspace.output_dir.iterdir())
    assert _output_rows(durable_workspace.session_factory) == []


@pytest.mark.asyncio
async def test_render_cleans_up_when_rasterization_fails_midway(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])

    def failing_raster(pdf_path, page_numbers, output_dir, **kwargs):
        output_dir = pathlib.Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "slide_001.png").write_bytes(PNG_1X1)
        raise render.PptxRenderRasterError("page 2 could not be rasterized")

    monkeypatch.setattr(render, "render_pdf_pages_to_pngs", failing_raster)
    monkeypatch.setattr(render, "_open_pdf", lambda _path: _FakeDocument([None] * 2))

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "raster_failed"
    assert not list(durable_workspace.workspace.output_dir.iterdir())
    assert _output_rows(durable_workspace.session_factory) == []


@pytest.mark.asyncio
async def test_render_cleans_up_when_a_png_cannot_be_written(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    events: list = []
    fake_fitz = _FakeFitz([(720.0, 540.0)], events)

    class _UnwritableExtent(_FakePixelExtent):
        def save(self, path) -> None:
            raise OSError("no space left on device")

    fake_fitz.pages[0].get_pixmap = lambda *, matrix, alpha: _UnwritableExtent(
        1440, 1080
    )
    monkeypatch.setitem(sys.modules, "fitz", fake_fitz)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is False
    assert result["error_code"] == "raster_failed"
    assert not list(durable_workspace.workspace.output_dir.iterdir())
    assert _output_rows(durable_workspace.session_factory) == []


@pytest.mark.asyncio
async def test_render_keeps_already_registered_pngs_when_registration_fails_partway(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=3, requested=[], rendered=[])
    workspace = durable_workspace.workspace
    real_register = workspace.register_delivery_file
    calls = 0

    def partial_registration(path):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("registration interrupted")
        return real_register(path)

    monkeypatch.setattr(workspace, "register_delivery_file", partial_registration)

    result = await render.render_pptx_to_images(str(deck), workspace)

    assert result["success"] is False
    assert result["error_code"] == "registration_failed"
    output_dirs = [path for path in workspace.output_dir.iterdir() if path.is_dir()]
    assert len(output_dirs) == 1
    remaining = sorted(path.name for path in output_dirs[0].iterdir())
    assert remaining == ["slide_001.png", "slide_002.png", "slide_003.png"]
    rows = _output_rows(durable_workspace.session_factory)
    assert [row.filename for row in rows] == ["slide_001.png"]
    assert result["registered_pages"] == {1: rows[0].file_id}
    assert result["output_dir"] == str(output_dirs[0])


@pytest.mark.asyncio
async def test_render_two_page_deck_with_real_soffice(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice (soffice) is not installed on this host")

    try:
        import pptx
    except (
        ImportError
    ) as exc:  # pragma: no cover - dependency is part of the backend image
        pytest.fail(
            "LibreOffice is available but python-pptx is not installed, so the "
            f"real conversion cannot run: {exc}"
        )
    deck_path = durable_workspace.workspace.input_dir / "synthetic-two-slides.pptx"
    presentation = pptx.Presentation()
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(72, 72, 400000, 800000)
        textbox.text_frame.text = f"synthetic slide {index + 1}"
    presentation.save(str(deck_path))
    deck_hash = hashlib.sha256(deck_path.read_bytes()).hexdigest()

    result = await render.render_pptx_to_images(
        str(deck_path), durable_workspace.workspace
    )

    assert result["success"] is True, result
    assert result["total_pages"] == 2
    assert result["page_numbers"] == [1, 2]
    assert len(result["pages"]) == 2
    for page in result["pages"]:
        path = pathlib.Path(page["file_ref"]["file_path"])
        assert path.read_bytes().startswith(PNG_MAGIC)
        assert page["width"] * page["height"] <= render.MAX_PAGE_PIXELS
        assert page["width"] > 0 and page["height"] > 0
        resolved = durable_workspace.workspace.resolve_file_id_detached(
            page["file_ref"]["file_id"]
        )
        assert resolved is not None
        assert resolved.read_bytes().startswith(PNG_MAGIC)
    image_hashes = [
        hashlib.sha256(
            pathlib.Path(page["file_ref"]["file_path"]).read_bytes()
        ).hexdigest()
        for page in result["pages"]
    ]
    assert len(set(image_hashes)) == 2, (
        "the two synthetic slides must render to different images"
    )
    assert [
        pathlib.Path(page["file_ref"]["file_path"]).name for page in result["pages"]
    ] == [
        "slide_001.png",
        "slide_002.png",
    ]
    assert hashlib.sha256(deck_path.read_bytes()).hexdigest() == deck_hash
    rows = _output_rows(durable_workspace.session_factory)
    assert len(rows) == 2
    assert len(list(durable_workspace.storage_root.rglob("*.hidden.preview.pdf"))) == 1
    assert not list(pathlib.Path(tempfile.gettempdir()).glob("xagent-pptx-profile-*"))


def _write_pdf_page(tmp_path, width: float, height: float, name: str):
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    try:
        page = document.new_page(width=width, height=height)
        page.insert_text((72, 72), f"{width:g}x{height:g}")
        pdf_path = tmp_path / name
        document.save(str(pdf_path))
    finally:
        document.close()
    return pdf_path


def test_real_rasterizer_default_scale_renders_tall_page(tmp_path):
    pytest.importorskip("fitz")
    render = _load_render_module()
    pdf_path = _write_pdf_page(tmp_path, 1430.0, 1670.0, "tall.pdf")
    output_dir = tmp_path / "pngs"
    output_dir.mkdir()

    total_pages, pages = _rasterize(render, pdf_path, [1], output_dir)

    assert total_pages == 1
    assert len(pages) == 1
    assert pages[0].path.read_bytes().startswith(PNG_MAGIC)
    assert pages[0].width * pages[0].height <= render.MAX_PAGE_PIXELS
    assert abs(pages[0].width / pages[0].height - 1430 / 1670) < 0.01


def test_real_rasterizer_default_scale_renders_wide_page(tmp_path):
    pytest.importorskip("fitz")
    render = _load_render_module()
    pdf_path = _write_pdf_page(tmp_path, 1810.0, 1080.0, "wide.pdf")
    output_dir = tmp_path / "pngs"
    output_dir.mkdir()

    total_pages, pages = _rasterize(render, pdf_path, [1], output_dir)

    assert total_pages == 1
    assert len(pages) == 1
    assert pages[0].path.read_bytes().startswith(PNG_MAGIC)
    assert pages[0].width * pages[0].height <= render.MAX_PAGE_PIXELS
    assert abs(pages[0].width / pages[0].height - 1810 / 1080) < 0.01


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("total_pages", "rendered_count", "truncated", "unrendered"),
    [(20, 20, False, 0), (21, 20, True, 1)],
)
async def test_render_default_page_boundaries(
    durable_workspace,
    monkeypatch,
    tmp_path,
    total_pages,
    rendered_count,
    truncated,
    unrendered,
):
    render = _load_render_module()
    deck = _write_deck(durable_workspace.workspace.input_dir)
    requested: list[list[int]] = []
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(
        monkeypatch,
        render,
        page_count=total_pages,
        requested=requested,
        rendered=[],
    )

    result = await render.render_pptx_to_images(str(deck), durable_workspace.workspace)

    assert result["success"] is True
    assert result["total_pages"] == total_pages
    assert result["rendered_pages"] == rendered_count
    assert result["truncated"] is truncated
    assert result["unrendered_pages"] == unrendered
    assert requested == [list(range(1, rendered_count + 1))]
    if truncated:
        assert result["unrendered_page_ranges"] == [[21, 21]]
    else:
        assert result["unrendered_page_ranges"] == []


def test_one_sided_explicit_ranges_cannot_escape_the_page_limit():
    render = _load_render_module()

    selection = render.resolve_page_selection(21, 1, None)
    assert len(selection.page_numbers) == 20
    assert selection.truncated
    with pytest.raises(render.PptxRenderArgumentError):
        render.resolve_page_selection(21, None, 21)
    assert list(render.resolve_page_selection(21, 21, None).page_numbers) == [21]
    assert list(render.resolve_page_selection(21, None, 20).page_numbers) == list(
        range(1, 21)
    )


@pytest.mark.asyncio
async def test_output_symlink_escape_is_refused_before_writing(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    outside = tmp_path / "outside"
    outside.mkdir()
    workspace.output_dir.rmdir()
    workspace.output_dir.symlink_to(outside, target_is_directory=True)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    requested: list[list[int]] = []
    rendered: list[list] = []
    _install_rasterizer(
        monkeypatch,
        render,
        page_count=2,
        requested=requested,
        rendered=rendered,
    )

    result = await render.render_pptx_to_images(str(deck), workspace)

    assert result["success"] is False
    assert result["error_code"] == "output_unavailable"
    assert rendered == []
    assert requested == []
    assert list(outside.rglob("*")) == []


@pytest.mark.asyncio
async def test_preexisting_output_directory_is_not_deleted_on_collision(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    existing = workspace.output_dir / "pptx_render_00000000000000000000000000000000"
    existing.mkdir()
    sentinel = existing / "existing.png"
    sentinel.write_bytes(b"pre-existing output")
    monkeypatch.setattr(
        render.uuid,
        "uuid4",
        lambda: SimpleNamespace(hex="0" * 32),
    )
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])

    result = await render.render_pptx_to_images(str(deck), workspace)

    assert sentinel.exists(), (
        "a failed mkdir must not delete a directory this call never created"
    )
    assert sentinel.read_bytes() == b"pre-existing output"
    assert result["success"] is False
    assert result["error_code"] == "output_unavailable"


@pytest.mark.asyncio
async def test_delegated_workspace_registration_matches_persisted_record(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    delegated = TaskWorkspace(
        id="synthetic_delegated_render",
        base_dir=str(durable_workspace.base_dir),
        db_task_id=4242,
    )
    deck = _write_deck(delegated.input_dir)
    rendered: list[list] = []
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(
        monkeypatch, render, page_count=2, requested=[], rendered=rendered
    )

    result = await render.render_pptx_to_images(str(deck), delegated)

    rows = _output_rows(durable_workspace.session_factory)
    originals = [page.path for batch in rendered for page in batch]
    assert rows and len(rows) == 2
    assert result["success"] is True, (
        "committed delegated outputs were classified as unregistered by "
        f"their old local paths: {result}"
    )
    assert all(path.exists() for path in originals), (
        "cleanup removed rendered PNGs whose records had committed under the "
        "parent task's canonical directory"
    )
    parent_output = (
        pathlib.Path(durable_workspace.base_dir) / "web_task_4242" / "output"
    )
    for row in rows:
        assert row.task_id == 4242
        assert pathlib.Path(str(row.storage_path)).is_relative_to(parent_output)
    fresh_parent = TaskWorkspace(
        id="web_task_4242", base_dir=str(durable_workspace.base_dir)
    )
    for page in result["pages"]:
        assert fresh_parent.resolve_file_id_detached(page["file_id"]) is not None


@pytest.mark.asyncio
@pytest.mark.parametrize("existing_image", [False, True])
async def test_replaced_render_directory_never_modifies_outside_files(
    durable_workspace, monkeypatch, tmp_path, existing_image
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "sentinel.txt").write_bytes(b"preserve this file")
    if existing_image:
        (outside / "slide_001.png").write_bytes(b"preserve this image")
    before = {p.name: p.read_bytes() for p in outside.iterdir()}
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    fake_fitz = _FakeFitz([(72, 54)], [])
    get_pixmap = fake_fitz.pages[0].get_pixmap

    def replace_before_write(**kwargs):
        directory = next(workspace.output_dir.glob("pptx_render_*"))
        directory.rename(workspace.output_dir / "moved_render")
        directory.symlink_to(outside, target_is_directory=True)
        return get_pixmap(**kwargs)

    fake_fitz.pages[0].get_pixmap = replace_before_write
    monkeypatch.setitem(sys.modules, "fitz", fake_fitz)
    result = await render.render_pptx_to_images(str(deck), workspace)

    assert {p.name: p.read_bytes() for p in outside.iterdir()} == before
    assert result["success"] is False, result
    assert (
        workspace.output_dir / "moved_render" / "slide_001.png"
    ).read_bytes() == PNG_1X1


@pytest.mark.asyncio
async def test_render_enables_hidden_slides_for_pdf_export(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    source = _write_deck(workspace.input_dir)
    options = []

    async def convert(source, output, **kwargs):
        options.append(kwargs)
        pdf = output / "deck.pdf"
        pdf.write_bytes(b"%PDF synthetic")
        return pdf

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert)
    _install_rasterizer(monkeypatch, render, page_count=3, requested=[], rendered=[])
    result = await render.render_pptx_to_images(str(source), workspace)
    assert result["success"] is True, result
    assert options[0].get("include_hidden_slides") is True


def _write_hidden_middle_deck(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    presentation = Presentation()
    for number, color in enumerate([(220, 30, 40), (20, 180, 60), (30, 60, 220)], 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*color)
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        text.text_frame.text = f"Source slide {number}"
        if number == 2:
            slide._element.set("show", "0")
    presentation.save(str(path))


@pytest.mark.asyncio
@pytest.mark.parametrize("page_number, color", [(2, (20, 180, 60)), (3, (30, 60, 220))])
async def test_real_soffice_preserves_hidden_source_slide_numbers(
    durable_workspace, page_number, color
):
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice is not installed")
    from PIL import Image

    render = _load_render_module()
    workspace = durable_workspace.workspace
    source = workspace.input_dir / "hidden-middle.pptx"
    _write_hidden_middle_deck(source)
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    result = await render.render_pptx_to_images(
        str(source), workspace, start_page=page_number, end_page=page_number
    )
    assert result["success"] is True, result
    assert result["total_pages"] == 3
    assert result["page_numbers"] == [page_number]
    rows = _output_rows(durable_workspace.session_factory)
    assert len(rows) == 1 and rows[0].storage_status == "available"
    with Image.open(rows[0].storage_path) as image:
        assert image.convert("RGB").getpixel((10, 10)) == color
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original


@pytest.mark.asyncio
async def test_real_soffice_preview_keeps_default_hidden_slide_behavior(tmp_path):
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice is not installed")
    import fitz

    conversion = _load_conversion_module()
    source = tmp_path / "hidden-middle.pptx"
    _write_hidden_middle_deck(source)
    output = tmp_path / "pdf"
    output.mkdir()
    pdf = await conversion.convert_pptx_to_pdf(source, output)
    with fitz.open(pdf) as document:
        assert document.page_count == 2
        assert "Source slide 3" in document[1].get_text()


@pytest.mark.asyncio
async def test_changed_input_symlink_is_rejected_before_conversion(
    durable_workspace, monkeypatch, tmp_path
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    outside = tmp_path / "private.pptx"
    outside.write_bytes(b"private presentation")
    resolve = render._resolve_pptx_input
    reads = []

    def resolve_then_replace(*args):
        source = resolve(*args)
        source.unlink()
        source.symlink_to(outside)
        return source

    async def convert(source, output_dir, **kwargs):
        reads.append(source.read_bytes())
        pdf = output_dir / "deck.pdf"
        pdf.write_bytes(b"%PDF synthetic")
        return pdf

    monkeypatch.setattr(render, "_resolve_pptx_input", resolve_then_replace)
    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert)
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])
    result = await render.render_pptx_to_images(str(deck), workspace)

    assert reads == [], "conversion read a substituted input"
    assert result["success"] is False
    assert not list(workspace.output_dir.iterdir())


@pytest.mark.asyncio
async def test_converter_reads_private_staged_input_and_discards_it(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    sources = []

    async def convert(source, output_dir, **kwargs):
        sources.append(source)
        assert source != deck
        assert not source.is_relative_to(workspace.workspace_dir)
        assert source.read_bytes() == deck.read_bytes()
        pdf = output_dir / "deck.pdf"
        pdf.write_bytes(b"%PDF synthetic")
        return pdf

    monkeypatch.setattr(render, "convert_pptx_to_pdf", convert)
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])
    result = await render.render_pptx_to_images(str(deck), workspace)

    assert result["success"] is True, result
    assert len(sources) == 1
    assert not sources[0].exists()


@pytest.mark.asyncio
async def test_registration_failure_keeps_all_written_images(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=3, requested=[], rendered=[])

    def reject(*args, **kwargs):
        raise RuntimeError("metadata database unavailable")

    monkeypatch.setattr(workspace, "register_files", reject)
    monkeypatch.setattr(workspace, "register_delivery_file", reject)
    result = await render.render_pptx_to_images(str(deck), workspace)

    assert result["error_code"] == "registration_failed"
    assert len(list(workspace.output_dir.glob("pptx_render_*/*.png"))) == 3


@pytest.mark.asyncio
async def test_cancellation_does_not_wait_for_registration(
    durable_workspace, monkeypatch
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    deck = _write_deck(workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=1, requested=[], rendered=[])
    ready = asyncio.Event()
    release = threading.Event()
    loop = asyncio.get_running_loop()

    def blocked(*args, **kwargs):
        loop.call_soon_threadsafe(ready.set)
        assert release.wait(5)
        raise RuntimeError("registration stopped")

    monkeypatch.setattr(workspace, "register_files", blocked)
    monkeypatch.setattr(workspace, "register_delivery_file", blocked)
    task = asyncio.create_task(render.render_pptx_to_images(str(deck), workspace))
    try:
        await asyncio.wait_for(ready.wait(), 3)
        task.cancel()
        done, _ = await asyncio.wait({task}, timeout=0.5)
        assert task in done, "cancellation waited for the registration worker"
        assert task.cancelled()
    finally:
        release.set()
        await asyncio.gather(task, return_exceptions=True)


@pytest.mark.parametrize("dimensions", [(0, 10), (10, 0), (float("inf"), 10)])
def test_rasterizer_rejects_invalid_page_dimensions(tmp_path, monkeypatch, dimensions):
    render = _load_render_module()
    monkeypatch.setitem(sys.modules, "fitz", _FakeFitz([dimensions], []))
    with pytest.raises(render.PptxRenderRasterError, match="dimensions"):
        _rasterize(render, tmp_path / "deck.pdf", [1], tmp_path)


@pytest.mark.parametrize("start,end", [(0, 2), (-1, 2), (2, 1), (1, 0)])
def test_page_range_rejects_invalid_bounds(start, end):
    render = _load_render_module()
    with pytest.raises(render.PptxRenderArgumentError):
        render.resolve_page_selection(5, start, end)


def test_cold_import_without_optional_renderer_keeps_tool_registry():
    code = """
import asyncio
import sys
sys.modules["fitz"] = None
from xagent.core.tools.adapters.vibe.factory import ToolRegistry
ToolRegistry._import_tool_modules()
assert ToolRegistry._modules_imported
modules = {creator.__module__.rsplit(".", 1)[-1] for creator, _, _ in ToolRegistry._tool_creators}
assert {"pptx_render_tool", "pptx_tool", "vision_tool", "workspace_file_tool"} <= modules
from xagent.core.tools.core.pptx_render import render_pptx_to_images
result = asyncio.run(render_pptx_to_images("deck.pptx", None))
assert result["error_code"] == "dependency_unavailable", result
assert "PyMuPDF" in result["error"]
"""
    completed = subprocess.run(
        [sys.executable, "-B", "-c", code], capture_output=True, text=True, timeout=30
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr


def test_start_only_range_is_capped_and_reports_truncation():
    render = _load_render_module()
    selection = render.resolve_page_selection(80, 21)
    assert selection.page_numbers == tuple(range(21, 41))
    assert selection.truncated is True
    assert render.resolve_page_selection(24, 21).page_numbers == (21, 22, 23, 24)
    assert render.resolve_page_selection(24, 21).truncated is False


@pytest.mark.asyncio
async def test_oversized_input_has_a_distinct_error(durable_workspace, monkeypatch):
    render = _load_render_module()
    source = _write_deck(durable_workspace.workspace.input_dir)
    monkeypatch.setattr(render, "get_max_upload_size_bytes", lambda: 1)
    result = await render.render_pptx_to_images(
        str(source), durable_workspace.workspace
    )
    assert result["error_code"] == "input_too_large", result


@pytest.mark.asyncio
@pytest.mark.parametrize("registered", [False, True])
async def test_render_cache_reuses_content_and_uses_existing_cleanup(
    durable_workspace, monkeypatch, registered
):
    from xagent.web.api.files import _pptx_pdf_cache_path
    from xagent.web.services.uploaded_file_cleanup_resources import capture_previews
    from xagent.web.services.uploaded_file_store import (
        delete_legacy_preview_caches,
        delete_pptx_pdf_cache,
    )

    render = _load_render_module()
    workspace = durable_workspace.workspace
    source = _write_deck(workspace.input_dir)
    file_id = _register_input_deck(workspace, source) if registered else None
    calls = []
    _install_conversion(monkeypatch, render, calls=calls, timeout_records=[])
    _install_rasterizer(monkeypatch, render, page_count=3, requested=[], rendered=[])
    preview = _pptx_pdf_cache_path(source, file_id)
    preview.write_bytes(b"default preview excludes hidden slides")
    for number in (2, 3):
        result = await render.render_pptx_to_images(
            str(source), workspace, start_page=number, end_page=number
        )
        assert result["success"], result
    assert len(calls) == 1
    cached = list(preview.parent.glob("*.hidden.preview.pdf"))
    assert len(cached) == 1
    assert preview.read_bytes() == b"default preview excludes hidden slides"
    source.write_bytes(source.read_bytes() + b" changed")
    result = await render.render_pptx_to_images(str(source), workspace, end_page=1)
    assert result["success"], result
    assert len(calls) == 2
    cached = list(preview.parent.glob("*.hidden.preview.pdf"))
    assert len(cached) == 2
    if registered:
        captured = capture_previews(
            {
                "roots": ["", "", str(durable_workspace.storage_root)],
                "file_id": file_id,
                "generation": "render-cache-test",
            }
        )
        assert len(captured) == 3
        delete_pptx_pdf_cache(file_id)
    else:
        delete_legacy_preview_caches(source.resolve())
    assert not preview.exists()
    assert all(not path.exists() for path in cached)


@pytest.mark.asyncio
@pytest.mark.parametrize("page_count", [1, 3])
async def test_cancelled_raster_stops_before_registration(
    durable_workspace, monkeypatch, page_count
):
    render = _load_render_module()
    workspace = durable_workspace.workspace
    source = _write_deck(workspace.input_dir)
    _install_conversion(monkeypatch, render, calls=[], timeout_records=[])
    events = []
    fake = _FakeFitz([(72, 54)] * page_count, events)
    monkeypatch.setitem(sys.modules, "fitz", fake)
    loop = asyncio.get_running_loop()
    started = asyncio.Event()
    finished = asyncio.Event()
    release = threading.Event()
    pixmap = fake.pages[0].get_pixmap
    worker = render._render_and_register

    def blocked(**kwargs):
        loop.call_soon_threadsafe(started.set)
        assert release.wait(5)
        return pixmap(**kwargs)

    def observed(*args, **kwargs):
        try:
            return worker(*args, **kwargs)
        finally:
            loop.call_soon_threadsafe(finished.set)

    fake.pages[0].get_pixmap = blocked
    monkeypatch.setattr(render, "_render_and_register", observed)
    task = asyncio.create_task(render.render_pptx_to_images(str(source), workspace))
    try:
        await asyncio.wait_for(started.wait(), 5)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(task, 1)
    finally:
        release.set()
    await asyncio.wait_for(finished.wait(), 5)
    assert _output_rows(durable_workspace.session_factory) == []
    assert list(workspace.output_dir.iterdir()) == []
    assert len(events) == 1


@pytest.mark.asyncio
async def test_real_soffice_cache_reuses_hidden_export(durable_workspace, monkeypatch):
    if shutil.which("soffice") is None:
        pytest.skip("LibreOffice is not installed")
    from PIL import Image

    render = _load_render_module()
    workspace = durable_workspace.workspace
    source = workspace.input_dir / "cached-hidden.pptx"
    _write_hidden_middle_deck(source)
    calls = []
    convert = render.convert_pptx_to_pdf

    async def observed(*args, **kwargs):
        calls.append(kwargs)
        return await convert(*args, **kwargs)

    monkeypatch.setattr(render, "convert_pptx_to_pdf", observed)
    for number, color in ((2, (20, 180, 60)), (3, (30, 60, 220))):
        result = await render.render_pptx_to_images(
            str(source), workspace, start_page=number, end_page=number
        )
        assert result["success"], result
        assert result["total_pages"] == 3
        with Image.open(result["pages"][0]["file_ref"]["file_path"]) as image:
            assert image.convert("RGB").getpixel((10, 10)) == color
    assert len(calls) == 1
    assert calls[0]["include_hidden_slides"] is True
    assert len(list(durable_workspace.storage_root.rglob("*.hidden.preview.pdf"))) == 1
