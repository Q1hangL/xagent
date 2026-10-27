"""Behavior tests for explicit read/unpack failures on invalid .pptx packages.

The fixtures are built with python-pptx so the tests do not need Node.js.
Valid presentations must keep their previous read and unpack results; a file
that is not a readable .pptx package must produce an explicit failure result
that names the file, the reason and the single-file scope of the error.
"""

import os
import re
import tempfile
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from xagent.core.tools.adapters.vibe.pptx_tool import PPTXGenerationTool
from xagent.core.tools.core.pptx_tool import (
    clean_pptx,
    pack_pptx,
    read_pptx,
    unpack_pptx,
)
from xagent.core.workspace import TaskWorkspace

INVALID_PACKAGE = "this is not a valid .pptx package"
ZIP_REASON = "its ZIP structure is missing or incomplete"
DAMAGED_HINT = "The file may be damaged or incompletely uploaded."
REEXPORT_HINT = "Please re-export or re-upload this file."
SINGLE_FILE_LIMIT = (
    "This error applies only to this file; other files must be read separately."
)
CONTENT_TYPES_PART = '"[Content_Types].xml"'
PRESENTATION_PART = '"ppt/presentation.xml"'
RELATIONSHIPS_PART = '"ppt/_rels/presentation.xml.rels"'
XML_REASON = "its XML content is not well-formed"
PER_FILE_SEPARATELY = "report each result separately"
PER_FILE_ONE_AT_A_TIME = "one at a time"
READ_OTHERS_UNAFFECTED = "one file failing does not mean the other files cannot be read"
UNPACK_OTHERS_UNAFFECTED = (
    "one file failing does not mean the other files cannot be unpacked"
)
# Advice that asks the user for an archive instead of the original file.
ARCHIVE_ADVICE = re.compile(
    r"(?i)\b(?:send|share|provide|attach|give|compress)\b[^.]{0,40}?"
    r"\b(?:zip|compressed|archive)\b"
)

# Baseline expectations recorded from the fixed baseline implementation on the
# same python-pptx generated deck, not derived from the new implementation.
EXPECTED_STRUCTURE = {
    "slide_count": 2,
    "slides": [
        {"index": 0, "filename": "slide1.xml", "hidden": False},
        {"index": 1, "filename": "slide2.xml", "hidden": False},
    ],
    "titles": ["First slide title", "Second slide title"],
    "success": True,
}
EXPECTED_TEXT = {
    "text": (
        "- First slide title\n- First slide body\n\n"
        "- Second slide title\n- Second slide body\n\n"
    ),
    "success": True,
}


def _make_valid_deck(path: Path) -> Path:
    """Create a two-slide valid .pptx with titles and body text."""
    presentation = Presentation()
    for title, body in (
        ("First slide title", "First slide body"),
        ("Second slide title", "Second slide body"),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = body
    presentation.save(str(path))
    return path


def _zip_members(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


def _repackage(
    source: Path,
    target: Path,
    *,
    drop: tuple[str, ...] = (),
    replace: dict[str, bytes] | None = None,
) -> Path:
    """Rewrite a package, dropping members and replacing member bytes."""
    replace = replace or {}
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w") as dst:
        for info in src.infolist():
            if info.filename in drop:
                continue
            data = replace.get(info.filename, src.read(info.filename))
            dst.writestr(info, data)
    return target


def _truncate(source: Path, target: Path, *, keep: int = 600) -> Path:
    """Cut the archive tail so the central directory is gone but PK remains."""
    target.write_bytes(source.read_bytes()[:keep])
    return target


def _assert_invalid_package(result: dict, file_name: str, reason: str) -> None:
    assert result.get("success") is False
    error = result.get("error")
    assert isinstance(error, str) and error
    assert f'"{file_name}"' in error
    assert INVALID_PACKAGE in error
    assert reason in error
    assert DAMAGED_HINT in error
    assert REEXPORT_HINT in error
    assert SINGLE_FILE_LIMIT in error


@pytest.fixture
def created_temp_dirs(monkeypatch: pytest.MonkeyPatch) -> list[Path]:
    """Record every temporary directory the reader creates."""
    created: list[Path] = []
    real_mkdtemp = tempfile.mkdtemp

    def tracking_mkdtemp(*args, **kwargs):
        directory = real_mkdtemp(*args, **kwargs)
        created.append(Path(directory))
        return directory

    monkeypatch.setattr(tempfile, "mkdtemp", tracking_mkdtemp)
    return created


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    return _make_valid_deck(tmp_path / "sample_deck.pptx")


# ---------------------------------------------------------------------------
# Valid presentations keep the previous results.
# ---------------------------------------------------------------------------


def test_valid_presentation_structure_is_unchanged(deck: Path) -> None:
    assert read_pptx(str(deck)) == EXPECTED_STRUCTURE


def test_valid_presentation_text_is_unchanged(deck: Path) -> None:
    assert read_pptx(str(deck), extract_text=True) == EXPECTED_TEXT


def test_valid_presentation_unpack_extracts_every_member(
    deck: Path, tmp_path: Path
) -> None:
    output_dir = tmp_path / "unpacked_deck"
    members = _zip_members(deck)

    result = unpack_pptx(str(deck), str(output_dir))

    assert result["success"] is True
    assert result["output_dir"] == str(output_dir)
    assert result["file_count"] == len(members)
    assert result["message"] == f"Extracted {len(members)} files to {output_dir}"
    unpacked = sorted(
        str(path.relative_to(output_dir))
        for path in output_dir.rglob("*")
        if path.is_file()
    )
    assert unpacked == sorted(members)


def test_valid_presentation_unpack_pack_round_trip(deck: Path, tmp_path: Path) -> None:
    unpacked = tmp_path / "unpacked_for_pack"
    repacked = tmp_path / "repacked_deck.pptx"

    assert unpack_pptx(str(deck), str(unpacked))["success"] is True
    pack_result = pack_pptx(str(unpacked), str(repacked))

    assert pack_result["success"] is True
    assert Path(pack_result["output_path"]) == repacked
    assert read_pptx(str(repacked)) == EXPECTED_STRUCTURE


def test_valid_presentation_clean_keeps_required_parts(
    deck: Path, tmp_path: Path
) -> None:
    unpacked = tmp_path / "unpacked_for_clean"
    assert unpack_pptx(str(deck), str(unpacked))["success"] is True

    result = clean_pptx(str(unpacked))

    assert result["success"] is True
    assert (unpacked / "[Content_Types].xml").is_file()
    assert (unpacked / "ppt" / "presentation.xml").is_file()


# ---------------------------------------------------------------------------
# A truncated ZIP must be reported as an incomplete package.
# ---------------------------------------------------------------------------


def test_truncated_package_fails_structured_read(deck: Path, tmp_path: Path) -> None:
    truncated = _truncate(deck, tmp_path / "truncated_deck.pptx")

    assert truncated.read_bytes()[:2] == b"PK"
    with pytest.raises(zipfile.BadZipFile):
        zipfile.ZipFile(truncated)

    _assert_invalid_package(
        read_pptx(str(truncated)), "truncated_deck.pptx", ZIP_REASON
    )


def test_truncated_package_fails_text_read(deck: Path, tmp_path: Path) -> None:
    truncated = _truncate(deck, tmp_path / "truncated_deck.pptx")

    _assert_invalid_package(
        read_pptx(str(truncated), extract_text=True),
        "truncated_deck.pptx",
        ZIP_REASON,
    )


def test_truncated_package_fails_unpack(deck: Path, tmp_path: Path) -> None:
    truncated = _truncate(deck, tmp_path / "truncated_deck.pptx")
    output_dir = tmp_path / "truncated_output"

    _assert_invalid_package(
        unpack_pptx(str(truncated), str(output_dir)),
        "truncated_deck.pptx",
        ZIP_REASON,
    )
    assert not output_dir.exists()


# ---------------------------------------------------------------------------
# Non-ZIP and empty inputs must not surface a bare "File is not a zip file".
# ---------------------------------------------------------------------------


def test_plain_text_input_is_not_a_bare_zip_error(tmp_path: Path) -> None:
    plain = tmp_path / "plain_notes.pptx"
    plain.write_bytes(b"this is plain text, not a package\n")

    result = read_pptx(str(plain))

    _assert_invalid_package(result, "plain_notes.pptx", ZIP_REASON)
    assert "File is not a zip file" not in result["error"]


def test_empty_input_is_not_a_bare_zip_error(tmp_path: Path) -> None:
    empty = tmp_path / "empty_input.pptx"
    empty.write_bytes(b"")

    result = read_pptx(str(empty))

    _assert_invalid_package(result, "empty_input.pptx", ZIP_REASON)
    assert "File is not a zip file" not in result["error"]


def test_plain_text_input_fails_unpack_with_a_clear_reason(tmp_path: Path) -> None:
    plain = tmp_path / "plain_notes.pptx"
    plain.write_bytes(b"this is plain text, not a package\n")

    result = unpack_pptx(str(plain), str(tmp_path / "plain_output"))

    _assert_invalid_package(result, "plain_notes.pptx", ZIP_REASON)
    assert "File is not a zip file" not in result["error"]


def test_empty_input_fails_unpack_with_a_clear_reason(tmp_path: Path) -> None:
    empty = tmp_path / "empty_input.pptx"
    empty.write_bytes(b"")

    result = unpack_pptx(str(empty), str(tmp_path / "empty_output"))

    _assert_invalid_package(result, "empty_input.pptx", ZIP_REASON)


# ---------------------------------------------------------------------------
# A valid ZIP still needs the required OOXML parts.
# ---------------------------------------------------------------------------


def _missing_part_deck(deck: Path, tmp_path: Path, part: str) -> Path:
    name = Path(part).name
    return _repackage(deck, tmp_path / f"missing_{name}.pptx", drop=(part,))


def test_missing_content_types_fails_structured_read(
    deck: Path, tmp_path: Path
) -> None:
    broken = _missing_part_deck(deck, tmp_path, "[Content_Types].xml")

    result = read_pptx(str(broken))

    _assert_invalid_package(result, broken.name, CONTENT_TYPES_PART)
    assert PRESENTATION_PART not in result["error"]


def test_missing_content_types_fails_text_read(deck: Path, tmp_path: Path) -> None:
    broken = _missing_part_deck(deck, tmp_path, "[Content_Types].xml")

    result = read_pptx(str(broken), extract_text=True)

    _assert_invalid_package(result, broken.name, CONTENT_TYPES_PART)
    assert PRESENTATION_PART not in result["error"]


def test_missing_content_types_fails_unpack(deck: Path, tmp_path: Path) -> None:
    broken = _missing_part_deck(deck, tmp_path, "[Content_Types].xml")
    output_dir = tmp_path / "missing_content_types_output"

    result = unpack_pptx(str(broken), str(output_dir))

    _assert_invalid_package(result, broken.name, CONTENT_TYPES_PART)
    assert not output_dir.exists()


def test_missing_presentation_part_fails_structured_read(
    deck: Path, tmp_path: Path
) -> None:
    broken = _missing_part_deck(deck, tmp_path, "ppt/presentation.xml")

    result = read_pptx(str(broken))

    _assert_invalid_package(result, broken.name, PRESENTATION_PART)
    assert "error" in result
    assert result.get("success") is not True


def test_missing_presentation_part_fails_text_read(deck: Path, tmp_path: Path) -> None:
    broken = _missing_part_deck(deck, tmp_path, "ppt/presentation.xml")

    result = read_pptx(str(broken), extract_text=True)

    _assert_invalid_package(result, broken.name, PRESENTATION_PART)
    assert "text" not in result


def test_missing_presentation_part_fails_unpack(deck: Path, tmp_path: Path) -> None:
    broken = _missing_part_deck(deck, tmp_path, "ppt/presentation.xml")
    output_dir = tmp_path / "missing_presentation_output"

    result = unpack_pptx(str(broken), str(output_dir))

    _assert_invalid_package(result, broken.name, PRESENTATION_PART)
    assert not output_dir.exists()


def test_missing_relationships_part_fails_structured_read(
    deck: Path, tmp_path: Path
) -> None:
    broken = _repackage(
        deck,
        tmp_path / "missing_relationships.pptx",
        drop=("ppt/_rels/presentation.xml.rels",),
    )

    result = read_pptx(str(broken))

    _assert_invalid_package(result, broken.name, RELATIONSHIPS_PART)
    assert result.get("success") is not True


def test_missing_relationships_part_fails_text_read(deck: Path, tmp_path: Path) -> None:
    broken = _repackage(
        deck,
        tmp_path / "missing_relationships.pptx",
        drop=("ppt/_rels/presentation.xml.rels",),
    )

    result = read_pptx(str(broken), extract_text=True)

    _assert_invalid_package(result, broken.name, RELATIONSHIPS_PART)
    assert result.get("success") is not True


def test_malformed_presentation_xml_fails_structured_read(
    deck: Path, tmp_path: Path
) -> None:
    broken = _repackage(
        deck,
        tmp_path / "malformed_presentation.pptx",
        replace={"ppt/presentation.xml": b"<not-well-formed"},
    )

    result = read_pptx(str(broken))

    _assert_invalid_package(result, broken.name, XML_REASON)
    assert result.get("success") is not True


def test_malformed_presentation_xml_fails_text_read(deck: Path, tmp_path: Path) -> None:
    broken = _repackage(
        deck,
        tmp_path / "malformed_presentation.pptx",
        replace={"ppt/presentation.xml": b"<not-well-formed"},
    )

    result = read_pptx(str(broken), extract_text=True)

    _assert_invalid_package(result, broken.name, XML_REASON)
    assert result.get("success") is not True


def test_archive_without_required_parts_fails_structured_read(
    tmp_path: Path,
) -> None:
    empty_zip = tmp_path / "empty_archive.pptx"
    with zipfile.ZipFile(empty_zip, "w") as archive:
        archive.writestr("readme.txt", "no package parts here")

    result = read_pptx(str(empty_zip))

    _assert_invalid_package(result, empty_zip.name, CONTENT_TYPES_PART)
    assert result.get("success") is not True


def test_archive_without_required_parts_fails_unpack(tmp_path: Path) -> None:
    empty_zip = tmp_path / "empty_archive.pptx"
    with zipfile.ZipFile(empty_zip, "w") as archive:
        archive.writestr("readme.txt", "no package parts here")

    result = unpack_pptx(str(empty_zip), str(tmp_path / "empty_archive_output"))

    _assert_invalid_package(result, empty_zip.name, CONTENT_TYPES_PART)
    assert result.get("success") is not True


# ---------------------------------------------------------------------------
# Workspace, absolute path and file_id resolution; OS errors stay OS errors.
# ---------------------------------------------------------------------------


@pytest.fixture
def workspace(tmp_path: Path) -> TaskWorkspace:
    return TaskWorkspace(id="pptx_read_errors", base_dir=str(tmp_path / "uploads"))


def test_relative_workspace_path_reads_a_valid_presentation(
    workspace: TaskWorkspace,
) -> None:
    _make_valid_deck(workspace.input_dir / "sample_deck.pptx")

    assert read_pptx("sample_deck.pptx", workspace=workspace) == EXPECTED_STRUCTURE


def test_relative_workspace_path_reports_an_invalid_presentation(
    workspace: TaskWorkspace,
) -> None:
    plain = workspace.input_dir / "broken_deck.pptx"
    plain.write_bytes(b"not a package")
    _make_valid_deck(workspace.input_dir / "other_deck.pptx")

    result = read_pptx("broken_deck.pptx", workspace=workspace)

    _assert_invalid_package(result, "broken_deck.pptx", ZIP_REASON)
    assert "other_deck.pptx" not in result["error"]
    assert str(workspace.input_dir) not in result["error"]


def test_absolute_path_reports_an_invalid_presentation(
    tmp_path: Path, workspace: TaskWorkspace
) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    result = read_pptx(str(plain), workspace=workspace)

    _assert_invalid_package(result, "broken_deck.pptx", ZIP_REASON)
    assert str(tmp_path) not in result["error"]


def test_file_id_resolution_reads_a_valid_presentation(
    workspace: TaskWorkspace,
) -> None:
    source = _make_valid_deck(workspace.input_dir / "sample_deck.pptx")
    file_id = workspace.register_file(str(source))

    assert read_pptx(file_id, workspace=workspace) == EXPECTED_STRUCTURE


def test_file_id_resolution_uses_the_real_file_name_in_the_error(
    workspace: TaskWorkspace,
) -> None:
    source = workspace.input_dir / "broken_deck.pptx"
    source.write_bytes(b"not a package")
    file_id = workspace.register_file(str(source))

    result = read_pptx(file_id, workspace=workspace)

    _assert_invalid_package(result, "broken_deck.pptx", ZIP_REASON)
    assert file_id not in result["error"]
    assert str(workspace.input_dir) not in result["error"]


def test_missing_file_raises_instead_of_reporting_corruption(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        read_pptx(str(tmp_path / "absent_deck.pptx"))
    with pytest.raises(FileNotFoundError):
        unpack_pptx(str(tmp_path / "absent_deck.pptx"), str(tmp_path / "output"))


def test_missing_workspace_file_raises_instead_of_reporting_corruption(
    workspace: TaskWorkspace,
) -> None:
    with pytest.raises(FileNotFoundError):
        read_pptx("absent_deck.pptx", workspace=workspace)


@pytest.mark.skipif(
    hasattr(os, "geteuid") and os.geteuid() == 0,
    reason="root can read a mode-000 file",
)
def test_permission_error_raises_instead_of_reporting_corruption(
    tmp_path: Path,
) -> None:
    locked = _make_valid_deck(tmp_path / "locked_deck.pptx")
    locked.chmod(0o000)
    try:
        with pytest.raises(PermissionError):
            read_pptx(str(locked))
        with pytest.raises(PermissionError):
            unpack_pptx(str(locked), str(tmp_path / "locked_output"))
    finally:
        locked.chmod(0o644)


def test_unpack_output_failure_raises_instead_of_reporting_corruption(
    deck: Path, tmp_path: Path
) -> None:
    occupied = tmp_path / "occupied_output"
    occupied.write_text("this path is a file")

    with pytest.raises(FileExistsError):
        unpack_pptx(str(deck), str(occupied))


# ---------------------------------------------------------------------------
# Multiple files: each result is reported on its own.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("order", [("broken", "valid"), ("valid", "broken")])
def test_bad_and_valid_files_are_reported_separately(
    deck: Path, tmp_path: Path, order: tuple[str, str]
) -> None:
    valid = tmp_path / "valid_deck.pptx"
    valid.write_bytes(deck.read_bytes())
    broken = _truncate(deck, tmp_path / "broken_deck.pptx")
    paths = {"valid": valid, "broken": broken}

    results = {kind: read_pptx(str(paths[kind])) for kind in order}

    broken_error = results["broken"]["error"]
    assert results["broken"]["success"] is False
    assert '"broken_deck.pptx"' in broken_error
    assert "valid_deck.pptx" not in broken_error
    assert results["valid"] == EXPECTED_STRUCTURE


@pytest.mark.parametrize("order", [("broken", "valid"), ("valid", "broken")])
def test_bad_and_valid_files_unpack_separately(
    deck: Path, tmp_path: Path, order: tuple[str, str]
) -> None:
    valid = tmp_path / "valid_deck.pptx"
    valid.write_bytes(deck.read_bytes())
    broken = _truncate(deck, tmp_path / "broken_deck.pptx")
    paths = {"valid": valid, "broken": broken}
    outputs = {
        "valid": tmp_path / "valid_output",
        "broken": tmp_path / "broken_output",
    }

    results = {
        kind: unpack_pptx(str(paths[kind]), str(outputs[kind])) for kind in order
    }

    assert results["broken"]["success"] is False
    assert '"broken_deck.pptx"' in results["broken"]["error"]
    assert "valid_deck.pptx" not in results["broken"]["error"]
    assert results["valid"]["success"] is True
    assert outputs["valid"].is_dir()
    assert not outputs["broken"].exists()


# ---------------------------------------------------------------------------
# Registered tool guidance: each file is read/unpacked and reported separately.
# ---------------------------------------------------------------------------


def _registered_descriptions() -> dict[str, str]:
    return {tool.name: tool.description for tool in PPTXGenerationTool().get_tools()}


def _assert_per_file_guidance(text: str, others_unaffected: str) -> None:
    assert PER_FILE_ONE_AT_A_TIME in text
    assert PER_FILE_SEPARATELY in text
    assert others_unaffected in text


def test_registered_read_description_documents_per_file_reporting() -> None:
    description = _registered_descriptions()["read_pptx"]

    _assert_per_file_guidance(description, READ_OTHERS_UNAFFECTED)


def test_registered_unpack_description_documents_per_file_reporting() -> None:
    description = _registered_descriptions()["unpack_pptx"]

    _assert_per_file_guidance(description, UNPACK_OTHERS_UNAFFECTED)


def test_read_pptx_docstring_documents_per_file_reporting() -> None:
    _assert_per_file_guidance(read_pptx.__doc__ or "", READ_OTHERS_UNAFFECTED)


def test_unpack_pptx_docstring_documents_per_file_reporting() -> None:
    _assert_per_file_guidance(unpack_pptx.__doc__ or "", UNPACK_OTHERS_UNAFFECTED)


# ---------------------------------------------------------------------------
# Temporary directories and failed unpack artifacts are cleaned up.
# ---------------------------------------------------------------------------


def test_successful_read_removes_its_temp_directory(
    deck: Path, created_temp_dirs: list[Path]
) -> None:
    assert read_pptx(str(deck)) == EXPECTED_STRUCTURE

    assert len(created_temp_dirs) == 1
    assert not created_temp_dirs[0].exists()


def test_successful_unpack_removes_its_temp_directory(
    deck: Path, tmp_path: Path, created_temp_dirs: list[Path]
) -> None:
    assert unpack_pptx(str(deck), str(tmp_path / "unpacked"))["success"] is True

    assert len(created_temp_dirs) == 1
    assert not created_temp_dirs[0].exists()


def test_rejected_package_creates_no_temp_directory(
    deck: Path, tmp_path: Path, created_temp_dirs: list[Path]
) -> None:
    truncated = _truncate(deck, tmp_path / "truncated_deck.pptx")

    assert read_pptx(str(truncated))["success"] is False
    assert unpack_pptx(str(truncated), str(tmp_path / "out"))["success"] is False

    assert created_temp_dirs == []


def test_read_parse_failure_removes_its_temp_directory(
    deck: Path, tmp_path: Path, created_temp_dirs: list[Path]
) -> None:
    broken = _repackage(
        deck,
        tmp_path / "malformed_presentation.pptx",
        replace={"ppt/presentation.xml": b"<not-well-formed"},
    )

    assert read_pptx(str(broken))["success"] is False

    assert len(created_temp_dirs) == 1
    assert not created_temp_dirs[0].exists()


def test_extraction_failure_removes_its_temp_directory(
    deck: Path, monkeypatch: pytest.MonkeyPatch, created_temp_dirs: list[Path]
) -> None:
    def failing_extractall(self, *args, **kwargs):
        raise zipfile.BadZipFile("bad compressed data in member")

    monkeypatch.setattr(zipfile.ZipFile, "extractall", failing_extractall)

    result = read_pptx(str(deck))

    _assert_invalid_package(result, "sample_deck.pptx", ZIP_REASON)
    assert len(created_temp_dirs) == 1
    assert not created_temp_dirs[0].exists()


def test_failed_unpack_leaves_no_output_directory(deck: Path, tmp_path: Path) -> None:
    broken = _missing_part_deck(deck, tmp_path, "[Content_Types].xml")
    output_dir = tmp_path / "failed_unpack_output"

    assert unpack_pptx(str(broken), str(output_dir))["success"] is False

    assert not output_dir.exists()


# ---------------------------------------------------------------------------
# Error text facets: real file name, single-file scope, re-export advice.
# ---------------------------------------------------------------------------


def test_read_error_names_the_input_file(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = read_pptx(str(plain))["error"]

    assert '"broken_deck.pptx"' in error
    assert str(tmp_path) not in error


def test_read_error_limits_the_failure_to_this_one_file(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = read_pptx(str(plain))["error"]

    assert SINGLE_FILE_LIMIT in error


def test_read_error_recommends_re_export_or_re_upload(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = read_pptx(str(plain))["error"]

    assert REEXPORT_HINT in error
    assert ARCHIVE_ADVICE.search(error) is None


def test_read_error_keeps_zip_as_a_technical_reason(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = read_pptx(str(plain))["error"]

    assert "ZIP" in error
    assert ZIP_REASON in error
    assert DAMAGED_HINT in error
    assert "may be damaged" in error


def test_unpack_error_names_the_input_file(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = unpack_pptx(str(plain), str(tmp_path / "out"))["error"]

    assert '"broken_deck.pptx"' in error


def test_unpack_error_limits_the_failure_to_this_one_file(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = unpack_pptx(str(plain), str(tmp_path / "out"))["error"]

    assert SINGLE_FILE_LIMIT in error


def test_unpack_error_recommends_re_export_or_re_upload(tmp_path: Path) -> None:
    plain = tmp_path / "broken_deck.pptx"
    plain.write_bytes(b"not a package")

    error = unpack_pptx(str(plain), str(tmp_path / "out"))["error"]

    assert REEXPORT_HINT in error
    assert ARCHIVE_ADVICE.search(error) is None
